"""
Generator for: 12 simulacao de compra.pptx
"""

import os
from PIL import Image, ImageDraw, ImageFont
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt

from slide_builder import (
    TEMP_DIR, FONT_HERO, FONT_TITLE, FONT_SUBTITLE, FONT_HEADING,
    FONT_BODY, FONT_BODY_BOLD, FONT_SMALL, FONT_SMALL_BOLD, FONT_CODE,
    C_DARK_BG, C_CARD_BG, C_CARD_BORDER, C_TEXT_LIGHT, C_TEXT_MUTED, C_WHITE,
    create_base_canvas, draw_header, draw_browser_chrome,
    draw_wp_sidebar, draw_arrow_pointer, draw_check_icon, draw_cross_icon,
    draw_bullet_dot, draw_star_badge
)

PRIMARY_GREEN = (16, 185, 129)     # Emerald 500
ACCENT_BLUE = (37, 99, 235)        # Blue 600
ACCENT_AMBER = (245, 158, 11)      # Amber 500

def create_slide_1():
    img, draw = create_base_canvas()
    draw.rounded_rectangle([120, 160, 520, 215], radius=12, fill=PRIMARY_GREEN)
    draw.text((140, 172), "MODULO 12 • EXPERIENCIA", font=FONT_HEADING, fill=(15, 23, 42))
    
    draw.text((120, 250), "Simulacao Completa de Compra", font=FONT_HERO, fill=C_WHITE)
    draw.text((120, 335), "& Gestao de Pedidos no WooCommerce", font=FONT_HERO, fill=PRIMARY_GREEN)
    
    draw.text((120, 440), "A Jornada Completa do Cliente da Vitrine ao Checkout, Pagamento Pix/Cartao e Rotina Operacional do Lojista", font=FONT_SUBTITLE, fill=C_TEXT_MUTED)
    
    features = [
        ("A Jornada do Cliente (UX)", "Do primeiro clique na vitrine ao carrinho, selecao de frete, pagamento sem atrito e tela de confirmacao."),
        ("Testes de Meios de Pagamento", "Como validar transacoes de Pix com QR Code dinamico, Cartao de Credito e comunicacao com gateway."),
        ("Painel Operacional do Lojista", "Como gerenciar pedidos, mudar status (Pendente para Processando e Concluido) e despachar com rastreio.")
    ]
    for i, (title, desc) in enumerate(features):
        x = 120 + i * 570
        y = 540
        draw.rounded_rectangle([x, y, x + 530, y + 260], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
        draw.ellipse([x + 30, y + 30, x + 80, y + 80], fill=PRIMARY_GREEN)
        draw.text((x + 46, y + 38), str(i+1), font=FONT_HEADING, fill=(15, 23, 42))
        draw.text((x + 95, y + 42), title, font=FONT_HEADING, fill=C_WHITE)
        
        words = desc.split()
        line = ""
        ty = y + 105
        for w in words:
            if len(line + " " + w) > 34:
                draw.text((x + 30, ty), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                ty += 28
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((x + 30, ty), line, font=FONT_BODY, fill=C_TEXT_MUTED)

    draw.rounded_rectangle([120, 850, 1800, 960], radius=12, fill=(6, 78, 59))
    draw.text((160, 880), "OBJETIVO DA AULA:", font=FONT_BODY_BOLD, fill=ACCENT_AMBER)
    draw.text((360, 880), "Realizar uma compra real de teste de ponta a ponta para validar toda a integracao antes de abrir a loja ao publico.", font=FONT_BODY, fill=C_WHITE)

    path = os.path.join(TEMP_DIR, "mod12_slide1.png")
    img.save(path, quality=95)
    return path

def create_slide_2():
    img, draw = create_base_canvas()
    draw_header(draw, "Funil de Vendas", "As 6 Etapas do Funil de Conversao no E-Commerce", "Como o consumidor progride em cada tela e onde ocorrem as principais taxas de perda", PRIMARY_GREEN)
    
    stages = [
        ("1. Vitrine / Home", "100% Visitantes", "Atracao visual, banners de promocoes, busca e vitrines de produtos em destaque.", (59, 130, 246)),
        ("2. Pagina do Produto", "~45% Engajamento", "Avaliacao de fotos, especificacoes tecnicas, parcelamento e calculo de frete.", (147, 51, 234)),
        ("3. Carrinho de Compras", "~20% Intencao", "Revisao dos itens, aplicacao de cupons de desconto e confirmacao de valores.", ACCENT_AMBER),
        ("4. Checkout de Pagamento", "~10% Decisao", "Preenchimento de endereco, dados fiscais (CPF) e selecao do metodo de frete.", (236, 72, 153)),
        ("5. Conclusao da Compra", "~2% a 4% Conversao", "Aprovacao do Pix ou Cartao e geracao do pedido confirmado.", PRIMARY_GREEN),
        ("6. Pos-Venda & Fidelizacao", "Retencao de Clientes", "Emails de rastreamento, entrega no prazo e incentivo a recompras.", (20, 184, 166))
    ]
    
    for i, (title, pct, desc, col) in enumerate(stages):
        x = 80 + i * 295
        y = 260
        draw.rounded_rectangle([x, y, x + 275, y + 680], radius=16, fill=C_CARD_BG, outline=col, width=2)
        draw.rounded_rectangle([x, y, x + 275, y + 80], radius=16, fill=col)
        draw.rectangle([x, y + 50, x + 275, y + 80], fill=col)
        draw.text((x + 18, y + 16), title, font=FONT_BODY_BOLD, fill=C_WHITE)
        draw.text((x + 18, y + 46), pct, font=FONT_SMALL, fill=(240, 240, 240))
        
        words = desc.split()
        line = ""
        ty = y + 110
        for w in words:
            if len(line + " " + w) > 20:
                draw.text((x + 18, ty), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                ty += 28
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((x + 18, ty), line, font=FONT_BODY, fill=C_TEXT_LIGHT)

    path = os.path.join(TEMP_DIR, "mod12_slide2.png")
    img.save(path, quality=95)
    return path

def create_slide_3():
    img, draw = create_base_canvas()
    draw_header(draw, "Etapa 1: Vitrine", "A Experiencia de Navegacao e Escolha no Catalogo", "Como o cliente descobre produtos atraves de categorias, filtros e barra de busca", PRIMARY_GREEN)
    
    # Browser Mockup
    draw_browser_chrome(draw, 120, 250, 1050, 750, title="Loja Virtual - Catalogo de Produtos", url="https://minhaloja.com.br/loja/?categoria=calcados")
    
    # Store Header
    draw.rounded_rectangle([150, 315, 1140, 365], radius=6, fill=(18, 24, 36))
    draw.text((170, 330), "MINHA LOJA", font=FONT_BODY_BOLD, fill=PRIMARY_GREEN)
    draw.text((320, 332), "Novidades   Roupas   Calcados   Acessorios   Ofertas", font=FONT_SMALL, fill=C_TEXT_LIGHT)
    draw.rounded_rectangle([860, 325, 1120, 355], radius=6, fill=(34, 43, 60))
    draw.text((880, 332), "O que você procura?  [Busca]", font=FONT_SMALL, fill=C_TEXT_MUTED)
    
    # Sidebar Filters
    draw.rounded_rectangle([150, 385, 360, 780], radius=8, fill=(24, 32, 47), outline=(56, 68, 88), width=1)
    draw.text((170, 405), "Filtros da Loja", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw.text((170, 445), "• Faixa de Preco:", font=FONT_SMALL_BOLD, fill=ACCENT_AMBER)
    draw.text((170, 475), "  R$ 50 — R$ 350", font=FONT_SMALL, fill=C_TEXT_LIGHT)
    draw.text((170, 520), "• Tamanho:", font=FONT_SMALL_BOLD, fill=ACCENT_AMBER)
    draw.text((170, 550), "  [X] 39   [X] 40   [ ] 41", font=FONT_SMALL, fill=C_TEXT_LIGHT)
    draw.text((170, 595), "• Cor:", font=FONT_SMALL_BOLD, fill=ACCENT_AMBER)
    draw.text((170, 625), "  Preto (12)  Azul (8)", font=FONT_SMALL, fill=C_TEXT_LIGHT)
    
    # Product Grid
    for i in range(2):
        for j in range(2):
            px = 390 + j * 370
            py = 385 + i * 280
            draw.rounded_rectangle([px, py, px + 345, py + 260], radius=8, fill=(24, 32, 47), outline=(56, 68, 88), width=1)
            draw.rounded_rectangle([px + 15, py + 15, px + 140, py + 140], radius=6, fill=PRIMARY_GREEN)
            draw.text((px + 45, py + 70), "[FOTO]", font=FONT_SMALL_BOLD, fill=(15, 23, 42))
            draw.text((px + 155, py + 25), f"Tenis Pro Modelo {i*2+j+1}", font=FONT_BODY_BOLD, fill=C_WHITE)
            draw.text((px + 155, py + 60), "R$ 249,90", font=FONT_HEADING, fill=PRIMARY_GREEN)
            draw.text((px + 155, py + 100), "6x de R$ 41,65 s/ juros", font=FONT_SMALL, fill=C_TEXT_MUTED)
            draw.rounded_rectangle([px + 15, py + 195, px + 330, py + 240], radius=6, fill=ACCENT_BLUE)
            draw.text((px + 110, py + 208), "VER DETALHES", font=FONT_SMALL_BOLD, fill=C_WHITE)

    # Right explanations
    draw.rounded_rectangle([1210, 250, 1840, 1000], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1245, 290), "Boas Praticas na Vitrine:", font=FONT_HEADING, fill=C_WHITE)
    
    tips = [
        ("Velocidade de Carregamento", "A vitrine deve carregar em menos de 2 segundos. Use imagens otimizadas e cache do LiteSpeed/WP Super Cache."),
        ("Barra de Pesquisa Eficiente", "A busca do WooCommerce deve encontrar produtos tanto pelo titulo quanto por sinonimos cadastrados nas tags."),
        ("Filtros Variaveis Ativos", "Permita ao cliente selecionar seu tamanho e ver apenas o que esta disponivel em estoque."),
        ("Etiquetas Visuais Promocionais", "Selos de 'Mais Vendido', 'Frete Gratis' ou 'Lancamento' aumentam a taxa de clique em ate 40%.")
    ]
    ty = 360
    for t_t, t_d in tips:
        draw.text((1245, ty), t_t, font=FONT_BODY_BOLD, fill=PRIMARY_GREEN)
        words = t_d.split()
        line = ""
        ly = ty + 28
        for w in words:
            if len(line + " " + w) > 40:
                draw.text((1245, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                ly += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1245, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
        ty += 140

    path = os.path.join(TEMP_DIR, "mod12_slide3.png")
    img.save(path, quality=95)
    return path

def create_slide_4():
    img, draw = create_base_canvas()
    draw_header(draw, "Etapa 2: Produto", "A Decisao de Compra: Selecionando Variacoes e Frete", "O momento exato em que o comprador escolhe opcoes e clica no botao CTA", PRIMARY_GREEN)
    
    # Left browser product
    draw_browser_chrome(draw, 120, 250, 1050, 750, title="Produto - Selecao de Opcoes", url="https://minhaloja.com.br/produto/tenis-pro-black")
    
    # Left photo
    draw.rounded_rectangle([160, 320, 520, 680], radius=12, fill=(34, 43, 60))
    draw.text((250, 490), "[FOTO TENIS PRETO]", font=FONT_BODY_BOLD, fill=C_TEXT_MUTED)
    
    # Right product details
    draw.text((560, 320), "Tenis Performance Running Pro", font=FONT_HEADING, fill=C_WHITE)
    draw.text((560, 365), "R$ 299,90", font=FONT_HERO, fill=PRIMARY_GREEN)
    
    # Variations selection
    draw.text((560, 445), "Selecione a Cor: [ PRETO ]", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw.ellipse([560, 480, 600, 520], fill=(20, 20, 20), outline=PRIMARY_GREEN, width=2)
    draw.ellipse([620, 480, 660, 520], fill=(240, 240, 240))
    draw.ellipse([680, 480, 720, 520], fill=(37, 99, 235))
    
    draw.text((560, 545), "Selecione o Tamanho:", font=FONT_BODY_BOLD, fill=C_WHITE)
    sizes = ["39", "40", "41", "42"]
    for i, sz in enumerate(sizes):
        active = (sz == "41")
        draw.rounded_rectangle([560 + i * 65, 580, 615 + i * 65, 630], radius=6, fill=PRIMARY_GREEN if active else (24, 32, 47), outline=PRIMARY_GREEN, width=1)
        draw.text((577 + i * 65, 595), sz, font=FONT_BODY_BOLD, fill=(15, 23, 42) if active else C_WHITE)

    # Quantity & Add to Cart button
    draw.rounded_rectangle([560, 660, 640, 720], radius=8, fill=(24, 32, 47), outline=(56, 68, 88), width=1)
    draw.text((592, 678), "1", font=FONT_HEADING, fill=C_WHITE)
    
    draw.rounded_rectangle([660, 660, 980, 720], radius=8, fill=PRIMARY_GREEN)
    draw.text((710, 678), "ADICIONAR AO CARRINHO", font=FONT_BODY_BOLD, fill=(15, 23, 42))
    
    # Notification popup
    draw.rounded_rectangle([560, 745, 1140, 805], radius=8, fill=(6, 78, 59), outline=PRIMARY_GREEN, width=1)
    draw_check_icon(draw, 585, 775, radius=12, bg=PRIMARY_GREEN, fg=(15, 23, 42))
    draw.text((615, 765), "'Tenis Performance (Tam 41, Preto)' foi adicionado ao seu carrinho!", font=FONT_SMALL_BOLD, fill=C_WHITE)

    # Right: UX triggers
    draw.rounded_rectangle([1210, 250, 1840, 1000], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1245, 290), "Gatilhos da Pagina de Produto:", font=FONT_HEADING, fill=C_WHITE)
    
    ux_triggers = [
        ("Feedback Imediato de Adicao", "O aviso visual confirma que o produto entrou no carrinho, permitindo ao comprador continuar comprando ou ir direto para o checkout."),
        ("Bloqueio de Variacao sem Estoque", "Se o tamanho 42 estivesse esgotado, o botao ficaria desabilitado ou com risco tachado, evitando frustracoes."),
        ("Troca Dinamica de Foto", "Ao clicar na cor azul, a imagem principal deve trocar imediatamente para o modelo azul."),
        ("Mini-Carrinho Lateral (Ajax Cart)", "Uma barra lateral deslizante que abre sem recarregar a pagina aumenta em 18% a velocidade de compra.")
    ]
    uy = 360
    for u_t, u_d in ux_triggers:
        draw.text((1245, uy), u_t, font=FONT_BODY_BOLD, fill=ACCENT_AMBER)
        words = u_d.split()
        line = ""
        ly = uy + 28
        for w in words:
            if len(line + " " + w) > 40:
                draw.text((1245, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                ly += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1245, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
        uy += 140

    path = os.path.join(TEMP_DIR, "mod12_slide4.png")
    img.save(path, quality=95)
    return path

def create_slide_5():
    img, draw = create_base_canvas()
    draw_header(draw, "Etapa 3: Carrinho", "Revisao do Carrinho e Aplicacao de Cupom", "Como a tela do carrinho valida quantidades, cupons e apresenta o resumo antes do pagamento", PRIMARY_GREEN)
    
    # Left browser cart
    draw_browser_chrome(draw, 120, 250, 1050, 750, title="Carrinho de Compras - WooCommerce", url="https://minhaloja.com.br/carrinho/")
    
    # Cart items table
    draw.rounded_rectangle([150, 315, 1140, 520], radius=8, fill=(24, 32, 47), outline=(56, 68, 88), width=1)
    draw.rectangle([150, 315, 1140, 360], fill=(18, 24, 36))
    draw.text((170, 330), "Produto", font=FONT_SMALL_BOLD, fill=C_TEXT_LIGHT)
    draw.text((560, 330), "Preco", font=FONT_SMALL_BOLD, fill=C_TEXT_LIGHT)
    draw.text((720, 330), "Quantidade", font=FONT_SMALL_BOLD, fill=C_TEXT_LIGHT)
    draw.text((950, 330), "Subtotal", font=FONT_SMALL_BOLD, fill=C_TEXT_LIGHT)
    
    # Row 1
    draw.text((170, 385), "Tenis Performance Running Pro", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw.text((170, 415), "Cor: Preto  |  Tamanho: 41", font=FONT_SMALL, fill=C_TEXT_MUTED)
    draw.text((560, 400), "R$ 299,90", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((740, 400), "[ -  1  + ]", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw.text((950, 400), "R$ 299,90", font=FONT_HEADING, fill=PRIMARY_GREEN)
    
    # Coupon code row
    draw.rounded_rectangle([150, 455, 450, 500], radius=6, fill=(34, 43, 60))
    draw.text((170, 468), "PRIMEIRACOMPRA", font=FONT_CODE, fill=ACCENT_AMBER)
    draw.rounded_rectangle([465, 455, 640, 500], radius=6, fill=ACCENT_BLUE)
    draw.text((495, 468), "Aplicar Cupom", font=FONT_SMALL_BOLD, fill=C_WHITE)
    
    # Cart Totals box
    draw.rounded_rectangle([680, 545, 1140, 800], radius=8, fill=(24, 32, 47), outline=PRIMARY_GREEN, width=1)
    draw.text((710, 565), "Total no Carrinho", font=FONT_HEADING, fill=C_WHITE)
    draw.text((710, 610), "Subtotal:", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((1010, 610), "R$ 299,90", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((710, 645), "Cupom (10% OFF):", font=FONT_BODY, fill=PRIMARY_GREEN)
    draw.text((1010, 645), "- R$ 29,99", font=FONT_BODY_BOLD, fill=PRIMARY_GREEN)
    draw.text((710, 680), "Frete (SEDEX):", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((1010, 680), "R$ 18,90", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.line([710, 715, 1110, 715], fill=(56, 68, 88), width=1)
    draw.text((710, 730), "TOTAL:", font=FONT_HEADING, fill=C_WHITE)
    draw.text((980, 730), "R$ 288,81", font=FONT_HERO, fill=PRIMARY_GREEN)
    
    draw.rounded_rectangle([680, 815, 1140, 865], radius=6, fill=PRIMARY_GREEN)
    draw.text((785, 830), "FECHAR PEDIDO (CHECKOUT)", font=FONT_BODY_BOLD, fill=(15, 23, 42))

    # Right: Coupon strategies
    draw.rounded_rectangle([1210, 250, 1840, 1000], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1245, 290), "Estrategias de Carrinho:", font=FONT_HEADING, fill=C_WHITE)
    
    cart_tips = [
        ("Como Criar Cupons no WooCommerce", "Acesse Marketing > Cupons no painel. Voce pode criar desconto percentual (ex: 10%), desconto fixo (ex: R$ 20) ou frete gratis."),
        ("Limite de Uso por Usuario", "Configure o cupom para funcionar apenas 1 vez por CPF ou email, evitando que o mesmo cliente reutilize indefinidamente."),
        ("Gatilho de Valor Minimo", "Exija um valor minimo (ex: 'Cupom valido para compras acima de R$ 150') para aumentar o ticket medio da loja."),
        ("Barra de Progresso de Frete Gratis", "Adicione um aviso dinâmico: 'Adicione mais R$ 11,19 para GANHAR FRETE GRATIS!'. E um dos maiores indutores de compra.")
    ]
    cy = 360
    for c_t, c_d in cart_tips:
        draw.text((1245, cy), c_t, font=FONT_BODY_BOLD, fill=PRIMARY_GREEN)
        words = c_d.split()
        line = ""
        ly = cy + 28
        for w in words:
            if len(line + " " + w) > 40:
                draw.text((1245, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                ly += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1245, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
        cy += 140

    path = os.path.join(TEMP_DIR, "mod12_slide5.png")
    img.save(path, quality=95)
    return path

def create_slide_6():
    img, draw = create_base_canvas()
    draw_header(draw, "Etapa 4: Checkout", "Preenchimento de Dados e Autocompletar de CEP", "O Checkout transparente de uma pagina que reduz friccao e erros de digitacao", PRIMARY_GREEN)
    
    # Left browser Checkout Form
    draw_browser_chrome(draw, 120, 250, 1050, 750, title="Finalizar Compra - Checkout Seguro", url="https://minhaloja.com.br/checkout/")
    
    draw.text((150, 315), "Detalhes de Faturamento e Entrega", font=FONT_HEADING, fill=C_WHITE)
    
    chk_fields = [
        ("Nome Completo:", "Carlos Alberto da Silva"),
        ("CPF (Obrigatorio para Emissao de NF):", "123.456.789-00"),
        ("Telefone Celular / WhatsApp:", "(11) 98765-4321"),
        ("Email para Envio do Rastreio:", "carlos.silva@email.com"),
        ("CEP de Entrega:", "01310-100  [Busca Automatica de Endereco]"),
        ("Endereco e Numero:", "Av. Paulista, 1000 - Apto 82"),
        ("Bairro, Cidade e Estado:", "Bela Vista - Sao Paulo / SP")
    ]
    cy = 360
    for label, val in chk_fields:
        draw.text((150, cy), label, font=FONT_SMALL_BOLD, fill=C_TEXT_LIGHT)
        draw.rounded_rectangle([150, cy + 24, 1140, cy + 62], radius=6, fill=(24, 32, 47), outline=(56, 68, 88), width=1)
        draw.text((170, cy + 32), val, font=FONT_BODY, fill=PRIMARY_GREEN if "Av." in val or "01310" in val else C_WHITE)
        cy += 74

    # Right: Best Practices Checkout
    draw.rounded_rectangle([1210, 250, 1840, 1000], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1245, 290), "Regras de Ouro do Checkout:", font=FONT_HEADING, fill=C_WHITE)
    
    rules = [
        ("Busca Automatica por CEP (ViaCEP)", "Ao digitar os 8 digitos do CEP, a rua, bairro, cidade e estado devem ser preenchidos sozinhos instantaneamente. Economiza tempo e evita erros de entrega."),
        ("Campo de CPF Validador", "No Brasil, toda transportadora e os Correios exigem CPF para emissao de Declaracao de Conteudo ou Nota Fiscal. O plugin 'Brazilian Market on WooCommerce' adiciona este campo."),
        ("Elimine Campos Desnecessarios", "Remova campos inuteis como 'Nome da Empresa', 'Telefone Fixo' ou 'Fax'. Cada campo a menos aumenta a conversao em 5%."),
        ("Checkout Sem Cadastro Obrigatorio", "Permita 'Comprar como Convidado'. O WooCommerce pode criar a conta do cliente silenciosamente em segundo plano.")
    ]
    ry = 360
    for r_t, r_d in rules:
        draw.text((1245, ry), r_t, font=FONT_BODY_BOLD, fill=ACCENT_AMBER)
        words = r_d.split()
        line = ""
        ly = ry + 28
        for w in words:
            if len(line + " " + w) > 40:
                draw.text((1245, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                ly += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1245, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
        ry += 140

    path = os.path.join(TEMP_DIR, "mod12_slide6.png")
    img.save(path, quality=95)
    return path

def create_slide_7():
    img, draw = create_base_canvas()
    draw_header(draw, "Etapa 5: Frete no Checkout", "Selecao do Metodo de Envio em Tempo Real", "O momento em que as transportadoras sao consultadas e exibem prazos e custos", PRIMARY_GREEN)
    
    # Left: Shipping selection card
    draw.rounded_rectangle([120, 260, 1050, 980], radius=16, fill=C_CARD_BG, outline=PRIMARY_GREEN, width=2)
    draw.text((160, 295), "Escolha o Tipo de Entrega:", font=FONT_HEADING, fill=C_WHITE)
    draw.text((160, 335), "Endereco de Entrega: Sao Paulo / SP (CEP 01310-100)", font=FONT_SMALL, fill=C_TEXT_MUTED)
    
    ship_options = [
        ("SEDEX (Correios)", "Entrega em ate 2 dias uteis", "R$ 18,90", True, PRIMARY_GREEN),
        ("Jadlog .Package", "Entrega em ate 3 dias uteis", "R$ 14,50", False, (56, 68, 88)),
        ("PAC (Correios)", "Entrega em ate 5 dias uteis", "R$ 12,40", False, (56, 68, 88)),
        ("Retirada na Loja", "Retirar em nosso balcao (Paulista)", "GRATIS", False, (56, 68, 88))
    ]
    sy = 390
    for name, prazo, val, selected, col in ship_options:
        draw.rounded_rectangle([160, sy, 1000, sy + 110], radius=12, fill=(24, 32, 47), outline=col, width=2 if selected else 1)
        draw.ellipse([185, sy + 40, 215, sy + 70], fill=PRIMARY_GREEN if selected else (34, 43, 60), outline=PRIMARY_GREEN, width=2)
        if selected:
            draw.ellipse([193, sy + 48, 207, sy + 62], fill=(15, 23, 42))
        draw.text((235, sy + 25), name, font=FONT_HEADING, fill=C_WHITE)
        draw.text((235, sy + 65), prazo, font=FONT_SMALL, fill=C_TEXT_MUTED)
        draw.text((820, sy + 38), val, font=FONT_HEADING, fill=PRIMARY_GREEN if "GRATIS" in val or selected else C_TEXT_LIGHT)
        sy += 130

    # Right explanations
    draw.rounded_rectangle([1120, 260, 1840, 1000], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1155, 290), "Como o WooCommerce Atualiza:", font=FONT_HEADING, fill=C_WHITE)
    
    flow = [
        ("Consulta em Segundo Plano (AJAX)", "Assim que o cliente digita o CEP, uma requisicao assincrona bate na API do Melhor Envio sem precisar recarregar a pagina inteira."),
        ("Atualizacao Automatica do Valor Total", "Ao clicar na bolinha do SEDEX (R$ 18,90), o valor total do pedido no canto direito atualiza instantaneamente para R$ 288,81."),
        ("Gravacao no Banco de Dados", "O servico escolhido (ex: Correios SEDEX) e salvo nos metadados do pedido para que o lojista saiba exatamente qual etiqueta comprar."),
        ("Opcao Padrao Pre-Selecionada", "O WooCommerce sempre marca previamente a opcao mais barata ou a mais rapida para facilitar o fluxo de clique do comprador.")
    ]
    fy = 360
    for f_t, f_d in flow:
        draw.text((1155, fy), f_t, font=FONT_BODY_BOLD, fill=ACCENT_BLUE)
        words = f_d.split()
        line = ""
        ly = fy + 28
        for w in words:
            if len(line + " " + w) > 42:
                draw.text((1155, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                ly += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1155, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
        fy += 140

    path = os.path.join(TEMP_DIR, "mod12_slide7.png")
    img.save(path, quality=95)
    return path

def create_slide_8():
    img, draw = create_base_canvas()
    draw_header(draw, "Etapa 6: Pagamento Pix", "Pagamento Instantaneo com Pix (Mercado Pago)", "A geracao do QR Code dinamico e do codigo Pix Copia e Cola para aprovacao imediata", PRIMARY_GREEN)
    
    # Left Mockup Pix Screen
    draw.rounded_rectangle([120, 260, 1050, 980], radius=16, fill=C_CARD_BG, outline=PRIMARY_GREEN, width=2)
    draw.text((160, 295), "Pague com Pix e Tenha Aprovacao Imediata", font=FONT_HEADING, fill=C_WHITE)
    draw.text((160, 335), "Pedido #1042  |  Valor: R$ 288,81", font=FONT_SMALL_BOLD, fill=PRIMARY_GREEN)
    
    # QR Code placeholder box
    draw.rounded_rectangle([160, 380, 480, 700], radius=12, fill=C_WHITE)
    draw.rectangle([190, 410, 450, 670], fill=(20, 20, 20))
    # Fake QR code pixel squares
    draw.rectangle([210, 430, 270, 490], fill=C_WHITE)
    draw.rectangle([225, 445, 255, 475], fill=(20, 20, 20))
    draw.rectangle([370, 430, 430, 490], fill=C_WHITE)
    draw.rectangle([385, 445, 415, 475], fill=(20, 20, 20))
    draw.rectangle([210, 590, 270, 650], fill=C_WHITE)
    draw.rectangle([225, 605, 255, 635], fill=(20, 20, 20))
    draw.rectangle([310, 530, 340, 560], fill=C_WHITE)
    
    # Instructions right of QR
    draw.text((520, 390), "1. Abra o app do seu banco ou carteira", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw.text((520, 420), "2. Escolha pagar via Pix com QR Code", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((520, 455), "3. Aponte a camera do seu celular", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((520, 490), "4. O pagamento e confirmado em 5 segundos!", font=FONT_BODY_BOLD, fill=PRIMARY_GREEN)
    
    draw.text((520, 545), "Ou use o Pix Copia e Cola:", font=FONT_BODY_BOLD, fill=ACCENT_AMBER)
    draw.rounded_rectangle([520, 580, 1000, 630], radius=6, fill=(24, 32, 47), outline=(56, 68, 88), width=1)
    draw.text((535, 595), "00020126580014br.gov.bcb.pix0136...", font=FONT_CODE, fill=C_TEXT_MUTED)
    draw.rounded_rectangle([520, 645, 740, 690], radius=6, fill=PRIMARY_GREEN)
    draw.text((545, 658), "COPIAR CODIGO PIX", font=FONT_SMALL_BOLD, fill=(15, 23, 42))

    # Timer box
    draw.rounded_rectangle([160, 730, 1000, 810], radius=8, fill=(22, 78, 99))
    draw.text((190, 755), "Tempo restante para pagamento: 29 minutos e 45 segundos", font=FONT_BODY_BOLD, fill=ACCENT_AMBER)
    draw.text((190, 780), "Apos esse prazo, o pedido e cancelado e o produto volta para o estoque.", font=FONT_SMALL, fill=C_TEXT_LIGHT)

    # Right advantages
    draw.rounded_rectangle([1120, 260, 1840, 1000], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1155, 290), "Por que o Pix e o Preferido?", font=FONT_HEADING, fill=C_WHITE)
    
    pix_pros = [
        ("Aprovacao Instantanea 24/7", "Nao precisa esperar compensacao bancaria (como no boleto de 1 a 3 dias). Aprovou na hora, o lojista ja pode embalar!"),
        ("Menor Taxa para o Lojista", "No Mercado Pago, o Pix tem taxa de apenas 0,99% (contra 3,5% a 5% do cartao de credito). Margem de lucro muito maior."),
        ("WebHook de Notificacao (IPN)", "Assim que o cliente paga no banco dele, o Mercado Pago avisa o WooCommerce em milissegundos e o status do pedido muda para 'Processando' automaticamente."),
        ("Zero Risco de Chargeback", "O Pix nao possui contestacao de compra por clonagem de cartao. Dinheiro 100% garantido na conta.")
    ]
    py = 360
    for p_t, p_d in pix_pros:
        draw.text((1155, py), p_t, font=FONT_BODY_BOLD, fill=PRIMARY_GREEN)
        words = p_d.split()
        line = ""
        ly = py + 28
        for w in words:
            if len(line + " " + w) > 42:
                draw.text((1155, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                ly += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1155, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
        py += 140

    path = os.path.join(TEMP_DIR, "mod12_slide8.png")
    img.save(path, quality=95)
    return path

def create_slide_9():
    img, draw = create_base_canvas()
    draw_header(draw, "Etapa 7: Cartao de Credito", "Pagamento com Cartao de Credito & Antifraude", "Checkout transparente com selecao de parcelas e analise de risco em tempo real", PRIMARY_GREEN)
    
    # Left browser Credit Card form
    draw.rounded_rectangle([120, 260, 1050, 980], radius=16, fill=C_CARD_BG, outline=PRIMARY_GREEN, width=2)
    draw.text((160, 295), "Pagamento com Cartao de Credito", font=FONT_HEADING, fill=C_WHITE)
    
    card_fields = [
        ("Numero do Cartao:", "4532 •••• •••• 8910  [Visa / Mastercard / Elo]"),
        ("Nome Impresso no Cartao:", "CARLOS A SILVA"),
        ("Data de Validade (MM/AA):", "08/29"),
        ("Codigo de Seguranca (CVV):", "382"),
        ("Opcoes de Parcelamento:", "6x de R$ 48,13 sem juros (ou ate 12x de R$ 26,80)")
    ]
    fy = 360
    for label, val in card_fields:
        draw.text((160, fy), label, font=FONT_SMALL_BOLD, fill=C_TEXT_LIGHT)
        draw.rounded_rectangle([160, fy + 24, 1000, fy + 65], radius=6, fill=(24, 32, 47), outline=(56, 68, 88), width=1)
        draw.text((180, fy + 35), val, font=FONT_BODY, fill=C_WHITE)
        fy += 78

    draw.rounded_rectangle([160, 770, 1000, 830], radius=8, fill=PRIMARY_GREEN)
    draw.text((450, 788), "FINALIZAR PAGAMENTO", font=FONT_HEADING, fill=(15, 23, 42))

    # Right: Antifraud mechanism
    draw.rounded_rectangle([1120, 260, 1840, 1000], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1155, 290), "Como Funciona o Antifraude:", font=FONT_HEADING, fill=C_WHITE)
    
    security = [
        ("Tokenizacao Segura", "Os numeros do cartao NUNCA passam pelo banco de dados do seu WordPress. Eles sao criptografados e enviados diretamente ao servidor seguro do Mercado Pago (PCI-DSS)."),
        ("Analise de Risco Instantanea", "O sistema avalia geolocalizacao, dispositivo do comprador, historico do CPF e velocidade de digitacao para aprovar a compra em segundos."),
        ("Status de Pagamento", "• 'Aprovado' (Status vira Processando)\n• 'Em Analise' (Aguardando verificacao manual)\n• 'Recusado' (Saldo insuficiente ou suspeita de fraude)"),
        ("Protecao contra Chargeback", "O Mercado Pago oferece programa de protecao ao vendedor para pedidos entregues com comprovante de postagem dos Correios.")
    ]
    sy = 360
    for s_t, s_d in security:
        draw.text((1155, sy), s_t, font=FONT_BODY_BOLD, fill=ACCENT_AMBER)
        words = s_d.split()
        line = ""
        ly = sy + 28
        for w in words:
            if len(line + " " + w) > 42:
                draw.text((1155, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                ly += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1155, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
        sy += 140

    path = os.path.join(TEMP_DIR, "mod12_slide9.png")
    img.save(path, quality=95)
    return path

def create_slide_10():
    img, draw = create_base_canvas()
    draw_header(draw, "Etapa 8: Confirmacao", "A Pagina de Pedido Recebido (Thank You Page)", "Como tranquilizar o comprador e informar os proximos passos da entrega", PRIMARY_GREEN)
    
    # Left browser Thank You page
    draw_browser_chrome(draw, 120, 250, 1050, 750, title="Pedido Confirmado - Minha Loja", url="https://minhaloja.com.br/checkout/order-received/1042/")
    
    # Success banner
    draw.rounded_rectangle([160, 320, 1140, 420], radius=12, fill=(6, 78, 59), outline=PRIMARY_GREEN, width=2)
    draw_check_icon(draw, 205, 370, radius=24, bg=PRIMARY_GREEN, fg=(15, 23, 42))
    draw.text((250, 345), "Obrigado! Seu pedido foi recebido com sucesso.", font=FONT_HEADING, fill=C_WHITE)
    draw.text((250, 385), "Enviamos todos os detalhes e comprovante para o email: carlos.silva@email.com", font=FONT_SMALL, fill=(240, 240, 240))
    
    # Order metadata row
    draw.rounded_rectangle([160, 450, 1140, 530], radius=8, fill=(24, 32, 47), outline=(56, 68, 88), width=1)
    draw.text((180, 465), "Numero do Pedido:", font=FONT_SMALL, fill=C_TEXT_MUTED)
    draw.text((180, 490), "#1042", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw.text((370, 465), "Data:", font=FONT_SMALL, fill=C_TEXT_MUTED)
    draw.text((370, 490), "22/10/2026", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw.text((540, 465), "Total Pago:", font=FONT_SMALL, fill=C_TEXT_MUTED)
    draw.text((540, 490), "R$ 288,81", font=FONT_BODY_BOLD, fill=PRIMARY_GREEN)
    draw.text((720, 465), "Metodo de Pagamento:", font=FONT_SMALL, fill=C_TEXT_MUTED)
    draw.text((720, 490), "Pix (Mercado Pago)", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw.text((950, 465), "Metodo de Envio:", font=FONT_SMALL, fill=C_TEXT_MUTED)
    draw.text((950, 490), "SEDEX (Ate 2 dias)", font=FONT_BODY_BOLD, fill=C_WHITE)
    
    # Items summary
    draw.text((160, 565), "Itens do Pedido:", font=FONT_HEADING, fill=C_WHITE)
    draw.text((160, 610), "1x Tenis Performance Running Pro (Preto / Tam 41) .................... R$ 269,91", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((160, 650), "Frete (SEDEX) ..................................................................................... R$ 18,90", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((160, 690), "Endereco de Entrega: Av. Paulista, 1000 - Apto 82, Sao Paulo - SP, 01310-100", font=FONT_BODY, fill=C_TEXT_MUTED)

    # Right: Customer reassurance
    draw.rounded_rectangle([1210, 250, 1840, 1000], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1245, 290), "O que o Cliente Espera Agora?", font=FONT_HEADING, fill=C_WHITE)
    
    expectations = [
        ("Previsao Exata de Entrega", "Informar a data limite (ex: 'Receba ate 26 de Outubro')."),
        ("Canal de Suporte Imediato", "Botoes de 'Falar no WhatsApp' com numero do suporte reduzem a ansiedade de compras online."),
        ("Acesso a Area 'Minha Conta'", "O cliente pode acompanhar o status do pedido logando com o email cadastrado."),
        ("Incentivo a Recompra", "Algumas lojas oferecem um cupom exclusivo para a proxima compra ja nesta pagina!")
    ]
    ey = 360
    for e_t, e_d in expectations:
        draw.text((1245, ey), e_t, font=FONT_BODY_BOLD, fill=PRIMARY_GREEN)
        words = e_d.split()
        line = ""
        ly = ey + 28
        for w in words:
            if len(line + " " + w) > 40:
                draw.text((1245, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                ly += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1245, ly), line, font=FONT_BODY, fill=C_TEXT_MUTED)
        ey += 140

    path = os.path.join(TEMP_DIR, "mod12_slide10.png")
    img.save(path, quality=95)
    return path

def create_slide_11():
    img, draw = create_base_canvas()
    draw_header(draw, "Etapa 9: Emails", "Os Emails Automaticos Transacionais do WooCommerce", "Como a loja se comunica automaticamente com cliente e lojista em cada etapa", PRIMARY_GREEN)
    
    emails = [
        ("1. Novo Pedido Recebido", "Destinatario: Lojista / Admin", "Avisa a sua equipe que entrou uma venda na loja para que comecem a separar o estoque.", (59, 130, 246)),
        ("2. Pedido em Processamento", "Destinatario: Cliente", "Confirma ao comprador que o pagamento Pix ou Cartao foi APROVADO com sucesso.", PRIMARY_GREEN),
        ("3. Pedido Concluido / Enviado", "Destinatario: Cliente", "Enviado quando o pacote e despachado na transportadora com o codigo de rastreamento.", (147, 51, 234)),
        ("4. Pedido Cancelado", "Destinatario: Cliente e Lojista", "Disparado caso o Pix expire sem pagamento ou o cartao de credito seja reprovado.", (239, 68, 68)),
        ("5. Redefinicao de Senha", "Destinatario: Cliente", "Link seguro para recuperacao de acesso a conta cadastrada na loja.", ACCENT_AMBER),
        ("6. Nota do Cliente Adicionada", "Destinatario: Cliente", "Mensagens manuais enviadas pelo lojista (ex: 'Seu pacote saiu para entrega').", (20, 184, 166))
    ]
    
    # 2 columns x 3 rows
    for i, (title, dest, desc, col) in enumerate(emails):
        col_idx = i % 2
        row_idx = i // 2
        x = 120 + col_idx * 840
        y = 260 + row_idx * 230
        
        draw.rounded_rectangle([x, y, x + 800, y + 200], radius=16, fill=C_CARD_BG, outline=col, width=2)
        draw.rounded_rectangle([x + 30, y + 30, x + 80, y + 80], radius=8, fill=col)
        draw.text((x + 46, y + 40), "@", font=FONT_HEADING, fill=C_WHITE)
        draw.text((x + 100, y + 30), title, font=FONT_HEADING, fill=C_WHITE)
        draw.text((x + 100, y + 68), dest, font=FONT_SMALL_BOLD, fill=col)
        
        words = desc.split()
        line = ""
        ty = y + 115
        for w in words:
            if len(line + " " + w) > 52:
                draw.text((x + 30, ty), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                ty += 26
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((x + 30, ty), line, font=FONT_BODY, fill=C_TEXT_LIGHT)

    path = os.path.join(TEMP_DIR, "mod12_slide11.png")
    img.save(path, quality=95)
    return path

def create_slide_12():
    img, draw = create_base_canvas()
    draw_header(draw, "Visao do Lojista", "Gerenciamento de Pedidos no Painel do WooCommerce", "Acesse WooCommerce > Pedidos para visualizar e controlar todas as vendas realizadas", PRIMARY_GREEN)
    
    # Left browser Admin Orders Table
    draw_browser_chrome(draw, 80, 250, 1100, 750, title="Pedidos - WooCommerce Admin", url="https://minhaloja.com.br/wp-admin/edit.php?post_type=shop_order")
    draw_wp_sidebar(draw, 82, 298, 200, 700, active_item="WooCommerce")
    
    # Top filters
    draw.text((310, 320), "Pedidos", font=FONT_HEADING, fill=C_WHITE)
    draw.text((310, 360), "Todos (14)  |  Processando (4)  |  Pendente (2)  |  Concluido (8)", font=FONT_SMALL, fill=PRIMARY_GREEN)
    
    # Orders table
    draw.rounded_rectangle([310, 400, 1140, 720], radius=8, fill=(24, 32, 47), outline=(56, 68, 88), width=1)
    draw.rectangle([310, 400, 1140, 440], fill=(18, 24, 36))
    draw.text((330, 412), "Pedido", font=FONT_SMALL_BOLD, fill=C_TEXT_LIGHT)
    draw.text((450, 412), "Data", font=FONT_SMALL_BOLD, fill=C_TEXT_LIGHT)
    draw.text((580, 412), "Status", font=FONT_SMALL_BOLD, fill=C_TEXT_LIGHT)
    draw.text((770, 412), "Itens / Destinatario", font=FONT_SMALL_BOLD, fill=C_TEXT_LIGHT)
    draw.text((1020, 412), "Total", font=FONT_SMALL_BOLD, fill=C_TEXT_LIGHT)
    
    orders = [
        ("#1042", "Ha 5 min", "Processando", PRIMARY_GREEN, "Carlos Silva (1 item)", "R$ 288,81"),
        ("#1041", "Ha 1 hora", "Processando", PRIMARY_GREEN, "Mariana Costa (2 itens)", "R$ 412,00"),
        ("#1040", "Ha 3 horas", "Pendente", ACCENT_AMBER, "Roberto Lima (Aguardando Pix)", "R$ 159,90"),
        ("#1039", "Ontem", "Concluido", (59, 130, 246), "Fernanda Souza (Despachado)", "R$ 89,90")
    ]
    oy = 450
    for num, dt, st, st_col, client, val in orders:
        draw.line([310, oy, 1140, oy], fill=(48, 58, 76), width=1)
        draw.text((330, oy + 15), num, font=FONT_BODY_BOLD, fill=C_WHITE)
        draw.text((450, oy + 15), dt, font=FONT_SMALL, fill=C_TEXT_MUTED)
        
        # Badge status
        draw.rounded_rectangle([580, oy + 10, 730, oy + 45], radius=6, fill=st_col)
        draw.text((600, oy + 16), st, font=FONT_SMALL_BOLD, fill=C_WHITE if st != "Processando" else (15, 23, 42))
        
        draw.text((770, oy + 15), client, font=FONT_SMALL, fill=C_TEXT_LIGHT)
        draw.text((1020, oy + 15), val, font=FONT_BODY_BOLD, fill=C_WHITE)
        oy += 65

    # Right: Order status lifecycle
    draw.rounded_rectangle([1220, 250, 1840, 1000], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1255, 290), "Ciclo de Vida do Pedido:", font=FONT_HEADING, fill=C_WHITE)
    
    lifecycle = [
        ("1. Pendente (Pending)", "Pedido criado no site, mas ainda aguardando confirmacao de pagamento do gateway ou geracao do Pix."),
        ("2. Processando (Processing)", "Pagamento aprovado com sucesso! A mercadoria esta pronta para ser separada e empacotada."),
        ("3. Concluido (Completed)", "O pacote foi despachado e o codigo de rastreamento foi enviado ao comprador."),
        ("4. Cancelado (Cancelled)", "O cliente nao pagou o Pix no prazo ou o pagamento de cartao foi reprovado pelo antifraude.")
    ]
    ly = 360
    for l_t, l_d in lifecycle:
        draw.text((1255, ly), l_t, font=FONT_BODY_BOLD, fill=PRIMARY_GREEN)
        words = l_d.split()
        line = ""
        ty = ly + 28
        for w in words:
            if len(line + " " + w) > 40:
                draw.text((1255, ty), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                ty += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1255, ty), line, font=FONT_BODY, fill=C_TEXT_MUTED)
        ly += 140

    path = os.path.join(TEMP_DIR, "mod12_slide12.png")
    img.save(path, quality=95)
    return path

def create_slide_13():
    img, draw = create_base_canvas()
    draw_header(draw, "Expedicao & Rastreio", "Despachando a Mercadoria e Inserindo Rastreio", "Como concluir o pedido, anexar o codigo dos Correios/Jadlog e avisar o cliente", PRIMARY_GREEN)
    
    # Left: Order detail admin mockup
    draw.rounded_rectangle([120, 260, 1050, 980], radius=16, fill=C_CARD_BG, outline=PRIMARY_GREEN, width=2)
    draw.text((160, 295), "Detalhes do Pedido #1042 — Carlos Silva", font=FONT_HEADING, fill=C_WHITE)
    
    # Status select
    draw.text((160, 350), "Status do Pedido: [ Concluido v ]", font=FONT_BODY_BOLD, fill=PRIMARY_GREEN)
    
    # Actions
    draw.text((160, 400), "Endereco de Entrega:", font=FONT_SMALL_BOLD, fill=C_TEXT_LIGHT)
    draw.text((160, 425), "Av. Paulista, 1000 - Apto 82, Bela Vista - Sao Paulo / SP, 01310-100", font=FONT_BODY, fill=C_WHITE)
    
    # Tracking Code Box
    draw.rounded_rectangle([160, 480, 1000, 600], radius=8, fill=(24, 32, 47), outline=ACCENT_AMBER, width=2)
    draw.text((190, 505), "CODIGO DE RASTREAMENTO (MELHOR ENVIO / CORREIOS):", font=FONT_SMALL_BOLD, fill=ACCENT_AMBER)
    draw.text((190, 540), "NL 123 456 789 BR  [ SEDEX ]", font=FONT_HERO, fill=C_WHITE)
    
    # Notes box
    draw.rounded_rectangle([160, 630, 1000, 830], radius=8, fill=(24, 32, 47), outline=(56, 68, 88), width=1)
    draw.text((190, 650), "Notas do Pedido (Historico de Eventos):", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw.text((190, 690), "• 22/10 14:32: Pagamento de R$ 288,81 aprovado via Pix Mercado Pago.", font=FONT_SMALL, fill=PRIMARY_GREEN)
    draw.text((190, 725), "• 22/10 15:10: Etiqueta emitida no Melhor Envio.", font=FONT_SMALL, fill=C_TEXT_LIGHT)
    draw.text((190, 760), "• 22/10 16:45: Pacote postado na agencia dos Correios (Rastreio: NL123456789BR).", font=FONT_SMALL, fill=ACCENT_AMBER)
    draw.text((190, 795), "• 22/10 16:46: Email de 'Pedido Concluido' enviado ao comprador.", font=FONT_SMALL, fill=PRIMARY_GREEN)

    # Right: Checklist
    draw.rounded_rectangle([1120, 260, 1840, 1000], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1155, 290), "Rotina Diaria de Expedicao:", font=FONT_HEADING, fill=C_WHITE)
    
    daily = [
        ("1. Imprimir Pedidos 'Processando'", "Faca uma triagem pela manha de todos os pedidos aprovados na noite anterior."),
        ("2. Picking & Packing (Coleta e Embalagem)", "Colete as pecas na prateleira pelo SKU e embale com plastico bolha e caixa adequada."),
        ("3. Gerar Etiquetas em Lote", "No plugin do Melhor Envio, selecione os pedidos e gere todas as etiquetas e declaracoes com 1 clique."),
        ("4. Colar Etiqueta e NF", "Fixe a etiqueta no topo e a Declaracao de Conteudo ou Nota Fiscal no verso da embalagem."),
        ("5. Despacho e Notificacao", "Poste no ponto de coleta. O plugin atualiza o status para 'Concluido' e dispara o codigo de rastreio para o WhatsApp e email do cliente!")
    ]
    dy = 350
    for d_t, d_d in daily:
        draw.text((1155, dy), d_t, font=FONT_BODY_BOLD, fill=ACCENT_BLUE)
        words = d_d.split()
        line = ""
        ty = dy + 28
        for w in words:
            if len(line + " " + w) > 42:
                draw.text((1155, ty), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                ty += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1155, ty), line, font=FONT_BODY, fill=C_TEXT_MUTED)
        dy += 115

    path = os.path.join(TEMP_DIR, "mod12_slide13.png")
    img.save(path, quality=95)
    return path

def create_slide_14():
    img, draw = create_base_canvas()
    draw_header(draw, "Recuperacao de Vendas", "Recuperacao de Carrinhos Abandonados e Pix Pendente", "Como resgatar de 15% a 30% das vendas que nao foram pagas na primeira tentativa", PRIMARY_GREEN)
    
    strategies = [
        ("1. Mensagem Automatica no WhatsApp", "Disparo apos 15 minutos do Pix nao pago: 'Ola Carlos, vimos que seu pedido #1042 esta aguardando pagamento. Segue o codigo Pix para concluir sua compra'.", (16, 185, 129)),
        ("2. Sequencia de Emails de Abandono", "Email 1 (apos 1 hora): 'Voce esqueceu algo no carrinho?'.\nEmail 2 (apos 24 horas): 'Seus itens estao acabando no estoque'.\nEmail 3 (apos 48 horas): 'Ganhe 5% OFF extra com o cupom VOLTA5'.", ACCENT_BLUE),
        ("3. Pop-up de Intencao de Saida (Exit Intent)", "Quando o usuario move o cursor do mouse em direcao ao botao de fechar a aba do navegador, exibe um cupom de desconto imediato.", ACCENT_AMBER),
        ("4. Remarketing no Instagram e Facebook", "Atraves do Pixel da Meta, mostre exatamente a foto da mochila que o usuario visitou no feed do Instagram dele.", (236, 72, 153))
    ]
    
    # 4 strategy cards
    for i, (title, desc, col) in enumerate(strategies):
        x = 120 + i * 420
        y = 260
        draw.rounded_rectangle([x, y, x + 395, y + 420], radius=16, fill=C_CARD_BG, outline=col, width=2)
        draw.rounded_rectangle([x, y, x + 395, y + 80], radius=16, fill=col)
        draw.rectangle([x, y + 50, x + 395, y + 80], fill=col)
        draw.text((x + 18, y + 25), title, font=FONT_BODY_BOLD, fill=C_WHITE if col != (16, 185, 129) else (15, 23, 42))
        
        words = desc.split()
        line = ""
        ty = y + 110
        for w in words:
            if len(line + " " + w) > 26:
                draw.text((x + 20, ty), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                ty += 28
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((x + 20, ty), line, font=FONT_BODY, fill=C_TEXT_LIGHT)

    # Bottom Stat banner
    draw.rounded_rectangle([120, 720, 1800, 980], radius=16, fill=(24, 32, 47), outline=PRIMARY_GREEN, width=2)
    draw.text((160, 750), "O PODER DA RECUPERACAO AUTOMATICA:", font=FONT_HEADING, fill=PRIMARY_GREEN)
    draw.text((160, 800), "• Em media, 7 em cada 10 carrinhos sao abandonados antes de pagar no Brasil.", font=FONT_BODY, fill=C_WHITE)
    draw.text((160, 840), "• Lojas que possuem recuperacao via WhatsApp resgatam entre 20% e 35% desses clientes.", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((160, 880), "• Isso significa faturar R$ 5.000 a R$ 15.000 a mais por mes sem gastar 1 centavo extra em anuncios!", font=FONT_BODY_BOLD, fill=ACCENT_AMBER)

    path = os.path.join(TEMP_DIR, "mod12_slide14.png")
    img.save(path, quality=95)
    return path

def create_slide_15():
    img, draw = create_base_canvas()
    draw_header(draw, "Conclusao do Curso", "Checklist Geral: Sua Loja Virtual 100% Pronta", "Parabens! Todos os modulos de infraestrutura, catalogo, pagamentos e frete estao concluidos", PRIMARY_GREEN)
    
    modules_done = [
        ("Modulo 06: WordPress Instalado", "Painel configurado, tema leve e profissional ativo e banco de dados conectado."),
        ("Modulo 08: WooCommerce Configurado", "Moeda em Real (R$), localizacao Brasil, paginas automaticas e seguranca SSL."),
        ("Modulo 09: Mercado Livre & Mercado Pago", "Gateway com Pix e Cartao ativo, catalogos vinculados e estoque integrado."),
        ("Modulo 10: Integracao de Frete Concluida", "Melhor Envio conectado, prazos ajustados, taxas de embalagem e cubagem pronta."),
        ("Modulo 11: Catalogo de Produtos no Ar", "Produtos simples e variaveis com fotos profissionais, SKU, SEO e estoque."),
        ("Modulo 12: Fluxo de Compra Testado", "Simulacao completa realizada: pedido gerado, pago, despachado e rastreado com sucesso!")
    ]
    
    # 2 columns x 3 rows
    for i, (m_title, m_desc) in enumerate(modules_done):
        col_idx = i % 2
        row_idx = i // 2
        x = 120 + col_idx * 840
        y = 260 + row_idx * 210
        
        draw.rounded_rectangle([x, y, x + 800, y + 185], radius=16, fill=C_CARD_BG, outline=PRIMARY_GREEN, width=2)
        draw_check_icon(draw, x + 40, y + 45, radius=18, bg=PRIMARY_GREEN, fg=(15, 23, 42))
        draw.text((x + 80, y + 30), m_title, font=FONT_HEADING, fill=C_WHITE)
        
        words = m_desc.split()
        line = ""
        ty = y + 85
        for w in words:
            if len(line + " " + w) > 50:
                draw.text((x + 80, ty), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                ty += 26
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((x + 80, ty), line, font=FONT_BODY, fill=C_TEXT_MUTED)

    # Big celebratory footer
    draw.rounded_rectangle([120, 900, 1800, 980], radius=12, fill=(6, 78, 59), outline=PRIMARY_GREEN, width=2)
    draw.text((150, 925), "SUA LOJA VIRTUAL ESTA PRONTA PARA VENDER EM TODO O BRASIL!", font=FONT_HEADING, fill=C_WHITE)

    path = os.path.join(TEMP_DIR, "mod12_slide15.png")
    img.save(path, quality=95)
    return path


SLIDE_BUILDERS = [
    (create_slide_1, "Abertura do Modulo 12: Simulacao Completa de Compra e Gestao de Pedidos. Apresentar os objetivos finais: percorrer o fluxo completo do ponto de vista do comprador e ensinar o lojista a gerenciar pedidos no painel."),
    (create_slide_2, "O Funil de Vendas: Explicar as 6 etapas do funil de conversao no e-commerce (Vitrine, Produto, Carrinho, Checkout, Conclusao e Pos-Venda)."),
    (create_slide_3, "Etapa 1: A Vitrine: Ensinar a importancia de filtros laterais funcionais (preco, tamanho, cor), busca interna e tempo de carregamento rapido."),
    (create_slide_4, "Etapa 2: A Pagina do Produto: Demonstrar a selecao de variacoes (cor/tamanho), a troca automatica de foto e o feedback visual de adicao ao carrinho."),
    (create_slide_5, "Etapa 3: O Carrinho de Compras: Explicar o funcionamento de cupons de desconto, alteracao de quantidade e incentivo de frete gratis."),
    (create_slide_6, "Etapa 4: O Checkout de Alta Conversao: Ensinar as boas praticas de formulario simples, autocompletar de endereco via CEP e obrigatoriedade do CPF para Nota Fiscal."),
    (create_slide_7, "Etapa 5: Selecao de Frete no Checkout: Mostrar a atualizacao em tempo real dos valores de frete via AJAX e o recalculo do total do pedido."),
    (create_slide_8, "Etapa 6: Pagamento via Pix: Apresentar o QR Code dinamico, Pix Copia e Cola, tempo limite e confirmacao automatica via WebHook do Mercado Pago."),
    (create_slide_9, "Etapa 7: Pagamento com Cartao: Explicar a tokenizacao segura, parcelamento com juros pelo gateway e protecao antifraude integrada."),
    (create_slide_10, "Etapa 8: Tela de Confirmacao: Apresentar a pagina de pedido recebido (Thank You Page), numero do pedido e canal de suporte."),
    (create_slide_11, "Etapa 9: Emails Automaticos: Detalhar os 6 tipos de emails transacionais disparados pelo WooCommerce para o cliente e lojista."),
    (create_slide_12, "Visao do Lojista: O Painel de Pedidos: Navegar em WooCommerce > Pedidos e explicar os status (Pendente, Processando, Concluido, Cancelado)."),
    (create_slide_13, "Expedicao e Rastreio: Roteiro pratico de separacao, impressao de etiquetas no Melhor Envio, despacho nos Correios e envio do rastreio ao cliente."),
    (create_slide_14, "Recuperacao de Vendas: Tecnicas de resgate de Pix pendente e carrinhos abandonados via WhatsApp e sequencia de emails."),
    (create_slide_15, "Conclusao do Curso: Revisao geral de todos os modulos integrados (WordPress, WooCommerce, Mercado Livre, Frete, Catalogo e Checkout) comemorando a conclusao da loja virtual.")
]

def build_presentation_12():
    print("Generating slides for 12 simulacao de compra.pptx...")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    for idx, (func, notes_text) in enumerate(SLIDE_BUILDERS):
        img_path = func()
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(img_path, 0, 0, width=Inches(13.333), height=Inches(7.5))
        
        # Add instructor notes
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = f"[AULA E-COMMERCE - MODULO 12 - SLIDE {idx+1}]\n{notes_text}"
        print(f"  Slide {idx+1}/{len(SLIDE_BUILDERS)} generated.")

    out_name = "12 simulacao de compra.pptx"
    prs.save(out_name)
    print(f"Presentation saved successfully as '{out_name}' ({os.path.getsize(out_name)} bytes).")

if __name__ == "__main__":
    build_presentation_12()
