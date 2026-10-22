"""
Generator for: 8 instalacao de woocomerce.pptx
"""

import os
from PIL import Image, ImageDraw, ImageFont
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt

from slide_builder import (
    TEMP_DIR, FONT_HERO, FONT_TITLE, FONT_SUBTITLE, FONT_HEADING,
    FONT_BODY, FONT_BODY_BOLD, FONT_SMALL, FONT_SMALL_BOLD, FONT_CODE,
    C_DARK_BG, C_CARD_BG, C_CARD_BORDER, C_TEXT_LIGHT, C_TEXT_MUTED, C_WHITE,
    ACCENTS, create_base_canvas, draw_header, draw_browser_chrome,
    draw_wp_sidebar, draw_arrow_pointer, draw_check_icon, draw_cross_icon,
    draw_bullet_dot, draw_star_badge
)

MOD_ACCENT = ACCENTS["woo"]["primary"]     # (127, 84, 179)
MOD_SECONDARY = ACCENTS["woo"]["secondary"] # (155, 114, 207)

def create_slide_1():
    img, draw = create_base_canvas()
    # Hero Title & Badge
    draw.rounded_rectangle([120, 160, 480, 215], radius=12, fill=MOD_ACCENT)
    draw.text((140, 172), "MODULO 08 • E-COMMERCE", font=FONT_HEADING, fill=C_WHITE)
    
    draw.text((120, 250), "Instalacao e Configuracao", font=FONT_HERO, fill=C_WHITE)
    draw.text((120, 335), "do WooCommerce no WordPress", font=FONT_HERO, fill=MOD_SECONDARY)
    
    draw.text((120, 440), "Transformando seu WordPress em uma Loja Virtual Completa, Escallavel e Pronta para Vendas", font=FONT_SUBTITLE, fill=C_TEXT_MUTED)
    
    # 3 Feature highlights
    features = [
        ("Plugin Aberto e Gratis", "Mais de 25% de todo e-commerce global roda no WooCommerce. Sem comissoes por venda."),
        ("Controle Total", "Voce e dono dos seus dados, produtos, clientes e personalizacao sem limitacoes."),
        ("Ecossistema Completo", "Integracao nativa com meios de pagamento nacionais (Pix, Cartao) e transportadoras brasileiras.")
    ]
    for i, (title, desc) in enumerate(features):
        x = 120 + i * 570
        y = 540
        draw.rounded_rectangle([x, y, x + 530, y + 260], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
        draw.ellipse([x + 30, y + 30, x + 80, y + 80], fill=MOD_ACCENT)
        draw.text((x + 46, y + 38), str(i+1), font=FONT_HEADING, fill=C_WHITE)
        draw.text((x + 100, y + 42), title, font=FONT_HEADING, fill=C_TEXT_LIGHT)
        
        words = desc.split()
        line = ""
        dy = y + 105
        for w in words:
            if len(line + " " + w) > 38:
                draw.text((x + 30, dy), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                dy += 28
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((x + 30, dy), line, font=FONT_BODY, fill=C_TEXT_MUTED)

    # Bottom bar
    draw.rounded_rectangle([120, 850, 1800, 930], radius=12, fill=(24, 32, 47), outline=MOD_ACCENT, width=1)
    draw_check_icon(draw, 150, 890, radius=14, bg=(16, 185, 129))
    draw.text((180, 875), "Pre-requisito: WordPress 6.x ou superior instalado e funcionando", font=FONT_BODY_BOLD, fill=(16, 185, 129))
    draw.text((1250, 875), "Professor / Instrutor: Curso Pratico de E-commerce", font=FONT_BODY, fill=C_TEXT_MUTED)

    path = os.path.join(TEMP_DIR, "mod8_slide1.png")
    img.save(path)
    return path

def create_slide_2():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • Fundamentos", "O que e o WooCommerce?", "A plataforma de comercio eletronico mais popular do mundo para WordPress", MOD_ACCENT)
    
    # Left column: Architecture diagram
    draw.rounded_rectangle([80, 260, 920, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((115, 290), "Arquitetura do E-commerce no WordPress", font=FONT_HEADING, fill=C_WHITE)
    
    layers = [
        ("4. SEUS PRODUTOS E VENDAS", "Catalogo, pedidos, clientes, relatorios", (16, 185, 129)),
        ("3. PLUGINS DE EXTENSAO", "Mercado Pago (Pix/Cartao), Melhor Envio, Correios", MOD_ACCENT),
        ("2. WOOCOMMERCE CORE", "Carrinho, Checkout, Gestao de Estoque, Moeda R$", MOD_SECONDARY),
        ("1. WORDPRESS + TEMA", "CMS base, paginas institucionais, layout visual", (34, 113, 177)),
        ("0. HOSPEDAGEM + PHP + MYSQL", "Servidor web, Apache/Nginx, banco de dados", (71, 85, 105))
    ]
    
    for i, (layer_title, layer_desc, color) in enumerate(layers):
        ly = 350 + i * 115
        draw.rounded_rectangle([115, ly, 885, ly + 95], radius=12, fill=(20, 26, 38), outline=color, width=2)
        draw.rectangle([115, ly, 130, ly + 95], fill=color)
        draw.text((145, ly + 18), layer_title, font=FONT_BODY_BOLD, fill=C_WHITE)
        draw.text((145, ly + 52), layer_desc, font=FONT_SMALL, fill=C_TEXT_MUTED)

    # Right column: Why choose WooCommerce
    draw.rounded_rectangle([960, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((995, 290), "Principais Beneficios e Vantagens", font=FONT_HEADING, fill=C_WHITE)
    
    benefits = [
        ("Codigo Aberto (Open Source)", "Sem mensalidades fixas para usar a plataforma. Voce so paga a sua hospedagem e as taxas de transacao dos meios de pagamento."),
        ("Propriedade Real dos Dados", "Ao contrario de plataformas fechadas (Shopify, Nuvemshop), no WooCommerce os dados dos clientes e do catalogo sao 100% seus."),
        ("Flexibilidade Ilimitada", "Permite vender produtos fisicos, produtos digitais para download, servicos, ingressos, assinaturas recorrentes e reservas."),
        ("Adaptacao Total ao Brasil", "Suporta perfeitamente CEP brasileiro, CPF/CNPJ no checkout, emissao de NF-e, Pix com aprovacao imediata e transportadoras nacionais.")
    ]
    
    for i, (b_title, b_desc) in enumerate(benefits):
        by = 350 + i * 145
        draw.rounded_rectangle([995, by, 1805, by + 125], radius=12, fill=(24, 32, 47), outline=(45, 55, 72), width=1)
        draw_check_icon(draw, 1030, by + 35, radius=14, bg=MOD_ACCENT)
        draw.text((1065, by + 22), b_title, font=FONT_BODY_BOLD, fill=C_TEXT_LIGHT)
        
        words = b_desc.split()
        line = ""
        dy = by + 58
        for w in words:
            if len(line + " " + w) > 60:
                draw.text((1065, dy), line, font=FONT_SMALL, fill=C_TEXT_MUTED)
                dy += 22
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1065, dy), line, font=FONT_SMALL, fill=C_TEXT_MUTED)

    path = os.path.join(TEMP_DIR, "mod8_slide2.png")
    img.save(path)
    return path

def create_slide_3():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • Pre-requisitos", "Requisitos Tecnicos do Servidor", "O que seu ambiente precisa ter antes de instalar o WooCommerce", MOD_ACCENT)
    
    reqs = [
        ("Versao do PHP", "PHP 7.4 ou superior (Recomendado: PHP 8.1 ou 8.2)", "Garante alta performance no calculo de taxas, processamento de carrinho e seguranca.", (16, 185, 129)),
        ("Banco de Dados", "MySQL 5.7+ ou MariaDB 10.4+", "Necessario para suportar indices rapidos de produtos, pedidos e metadados.", (34, 113, 177)),
        ("Certificado SSL (HTTPS)", "Obrigatorio para E-commerce", "Sem HTTPS, navegadores alertam 'Site Nao Seguro' e gateways de pagamento bloqueiam o checkout.", (239, 68, 68)),
        ("Limite de Memoria PHP", "Minimo: 256 MB (Recomendado: 512 MB)", "Evita o erro 'Fatal Error: Allowed memory size exhausted' durante operacoes pesadas.", (245, 158, 11)),
        ("Extensoes PHP Ativas", "cURL, mbstring, OpenSSL, Zip, GD/Imagick", "Essenciais para comunicacao com APIs dos Correios, Mercado Pago e tratamento de fotos.", MOD_ACCENT)
    ]
    
    for i, (title, spec, why, col) in enumerate(reqs):
        y = 260 + i * 138
        draw.rounded_rectangle([80, y, 1840, y + 118], radius=14, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
        draw.rectangle([80, y, 92, y + 118], fill=col)
        
        draw.text((120, y + 20), title, font=FONT_HEADING, fill=C_WHITE)
        draw.rounded_rectangle([520, y + 20, 1100, y + 60], radius=8, fill=(15, 23, 42), outline=col, width=1)
        draw.text((535, y + 26), spec, font=FONT_BODY_BOLD, fill=col)
        
        draw.text((120, y + 72), "Motivo: " + why, font=FONT_BODY, fill=C_TEXT_MUTED)

    path = os.path.join(TEMP_DIR, "mod8_slide3.png")
    img.save(path)
    return path

def create_slide_4():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • Passo a Passo", "Passo 1: Acessar Menu de Plugins no WordPress", "Como navegar ate o gerenciador de extensoes no painel administrativo", MOD_ACCENT)
    
    draw_browser_chrome(draw, 80, 260, 1250, 720, "Painel Administrativo WordPress", "https://meusite.com/wp-admin/plugins.php")
    draw_wp_sidebar(draw, 80, 308, 250, 672, active_item="Plugins")
    
    draw.rectangle([330, 630, 520, 720], fill=(44, 51, 56))
    draw.text((345, 642), "Plugins Instalados", font=FONT_SMALL, fill=(200, 210, 220))
    draw.text((345, 680), "Adicionar Plugin", font=FONT_SMALL_BOLD, fill=(114, 168, 255))
    
    draw.rectangle([330, 308, 1330, 630], fill=(240, 242, 245))
    draw.text((370, 345), "Plugins", font=FONT_TITLE, fill=(30, 40, 55))
    draw.rounded_rectangle([520, 350, 720, 395], radius=6, fill=(34, 113, 177))
    draw.text((540, 360), "Adicionar Novo", font=FONT_BODY_BOLD, fill=C_WHITE)
    
    draw_arrow_pointer(draw, (850, 372), (730, 372), color=(239, 68, 68), width=6)
    
    # Right instruction card
    draw.rounded_rectangle([1370, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1405, 300), "Instrucoes do Passo 1", font=FONT_HEADING, fill=C_WHITE)
    
    instructions = [
        "1. Faca login no seu WordPress com a conta de Administrador.",
        "2. No menu lateral preto a esquerda, localize a opcao 'Plugins'.",
        "3. Passe o mouse sobre 'Plugins' e clique em 'Adicionar Plugin' (ou 'Adicionar Novo').",
        "4. Voce sera direcionado para o repositorio oficial com mais de 60.000 plugins gratuitos do WordPress.org."
    ]
    dy = 370
    for inst in instructions:
        words = inst.split()
        line = ""
        for w in words:
            if len(line + " " + w) > 30:
                draw.text((1405, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                dy += 30
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1405, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
            dy += 45

    # Tip box
    draw.rounded_rectangle([1400, 750, 1810, 930], radius=12, fill=(20, 35, 30), outline=(16, 185, 129), width=2)
    draw.text((1420, 770), "[DICA IMPORTANTE]", font=FONT_BODY_BOLD, fill=(16, 185, 129))
    draw.text((1420, 810), "Sempre utilize o repositorio", font=FONT_SMALL, fill=C_TEXT_LIGHT)
    draw.text((1420, 835), "oficial para garantir arquivos", font=FONT_SMALL, fill=C_TEXT_LIGHT)
    draw.text((1420, 860), "limpos e sem codigos maliciosos.", font=FONT_SMALL, fill=C_TEXT_LIGHT)

    path = os.path.join(TEMP_DIR, "mod8_slide4.png")
    img.save(path)
    return path

def create_slide_5():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • Passo a Passo", "Passo 2: Pesquisar o WooCommerce", "Localizando o plugin oficial no catalogo de plugins do WordPress", MOD_ACCENT)
    
    draw_browser_chrome(draw, 80, 260, 1250, 720, "Adicionar Plugins - WordPress", "https://meusite.com/wp-admin/plugin-install.php")
    draw_wp_sidebar(draw, 80, 308, 250, 672, active_item="Plugins")
    
    draw.rectangle([330, 308, 1330, 980], fill=(245, 247, 250))
    draw.text((370, 340), "Adicionar Plugins", font=FONT_TITLE, fill=(30, 40, 55))
    
    draw.rounded_rectangle([850, 340, 1280, 385], radius=6, fill=C_WHITE, outline=(180, 190, 205), width=2)
    draw.text((865, 348), "Pesquisar: WooCommerce", font=FONT_BODY_BOLD, fill=(40, 50, 70))
    draw_arrow_pointer(draw, (1100, 430), (1050, 390), color=(239, 68, 68), width=5)
    
    card_x, card_y = 370, 430
    draw.rounded_rectangle([card_x, card_y, card_x + 880, card_y + 450], radius=12, fill=C_WHITE, outline=(210, 220, 230), width=2)
    
    draw.rounded_rectangle([card_x + 30, card_y + 30, card_x + 120, card_y + 120], radius=16, fill=MOD_ACCENT)
    draw.text((card_x + 48, card_y + 45), "W", font=FONT_HERO, fill=C_WHITE)
    
    draw.text((card_x + 150, card_y + 35), "WooCommerce", font=FONT_HEADING, fill=(20, 30, 45))
    draw.text((card_x + 150, card_y + 75), "Por Automattic", font=FONT_BODY, fill=(100, 115, 130))
    
    desc_woo = "O WooCommerce e a plataforma de comercio eletronico de codigo aberto mais popular do mundo. Solucao principal gratuita, flexivel e ampliada por uma comunidade global."
    draw.text((card_x + 30, card_y + 145), desc_woo[:80], font=FONT_BODY, fill=(70, 80, 95))
    draw.text((card_x + 30, card_y + 180), desc_woo[80:], font=FONT_BODY, fill=(70, 80, 95))
    
    draw_star_badge(draw, card_x + 30, card_y + 250, count=5)
    draw.text((card_x + 150, card_y + 248), "(Mais de 5 milhoes de instalacoes ativas)", font=FONT_BODY_BOLD, fill=(234, 150, 20))
    
    draw_check_icon(draw, card_x + 45, card_y + 298, radius=12, bg=(16, 185, 129))
    draw.text((card_x + 65, card_y + 288), "Compativel com a sua versao do WordPress", font=FONT_BODY_BOLD, fill=(16, 185, 129))
    draw.text((card_x + 30, card_y + 325), "Ultima atualizacao: ha poucas semanas", font=FONT_SMALL, fill=(120, 130, 145))
    
    draw.rounded_rectangle([card_x + 650, card_y + 35, card_x + 840, card_y + 85], radius=6, fill=(34, 113, 177))
    draw.text((card_x + 685, card_y + 48), "Instalar Agora", font=FONT_BODY_BOLD, fill=C_WHITE)
    
    # Right column checklist
    draw.rounded_rectangle([1370, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1405, 300), "Criterios de Checagem", font=FONT_HEADING, fill=C_WHITE)
    
    items = [
        ("Autor Oficial", "Verifique se o autor e 'Automattic'. Existem dezenas de complementos, mas o principal e este."),
        ("Instalacoes Ativas", "Deve constar 'Mais de 5 milhoes de instalacoes ativas'."),
        ("Compatibilidade", "Deve exibir um selo verde confirmando compatibilidade com o seu WordPress."),
        ("Avaliacoes", "Milhares de notas 5 estrelas atestam a estabilidade do codigo.")
    ]
    dy = 370
    for t, d in items:
        draw_bullet_dot(draw, 1415, dy + 10, radius=5, color=MOD_SECONDARY)
        draw.text((1435, dy), t, font=FONT_BODY_BOLD, fill=MOD_SECONDARY)
        dy += 28
        words = d.split()
        line = ""
        for w in words:
            if len(line + " " + w) > 30:
                draw.text((1435, dy), line, font=FONT_SMALL, fill=C_TEXT_MUTED)
                dy += 22
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1435, dy), line, font=FONT_SMALL, fill=C_TEXT_MUTED)
            dy += 35

    path = os.path.join(TEMP_DIR, "mod8_slide5.png")
    img.save(path)
    return path

def create_slide_6():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • Passo a Passo", "Passo 3: Instalar e Ativar", "Executando o download dos arquivos e inicializando o WooCommerce", MOD_ACCENT)
    
    states = [
        ("ESTADO 1: DISPONIVEL", "O plugin foi encontrado no repositorio", "Instalar Agora", (34, 113, 177), "Clique no botao azul 'Instalar Agora'. O WordPress baixa o arquivo .zip diretamente do repositorio oficial para seu servidor."),
        ("ESTADO 2: PROCESSANDO", "Download e descompactacao dos arquivos", "Instalando...", (100, 115, 130), "O servidor descompacta os arquivos na pasta /wp-content/plugins/woocommerce/. Aguarde alguns segundos sem atualizar a pagina."),
        ("ESTADO 3: PRONTO PARA ATIVAR", "Arquivos instalados com sucesso", "Ativar", (16, 185, 129), "O botao mudara automaticamente para 'Ativar'. Clique nele para registrar as tabelas no MySQL e ativar os recursos de loja.")
    ]
    
    for i, (st_name, st_sub, btn_txt, btn_col, explanation) in enumerate(states):
        y = 260 + i * 235
        draw.rounded_rectangle([80, y, 1840, y + 205], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
        
        draw.ellipse([120, y + 30, 190, y + 100], fill=MOD_ACCENT)
        draw.text((145, y + 45), str(i+1), font=FONT_TITLE, fill=C_WHITE)
        
        draw.text((220, y + 30), st_name, font=FONT_HEADING, fill=C_WHITE)
        draw.text((220, y + 70), st_sub, font=FONT_BODY, fill=C_TEXT_MUTED)
        
        bx, by = 850, y + 35
        draw.rounded_rectangle([bx, by, bx + 220, by + 60], radius=8, fill=btn_col)
        draw.text((bx + 35, by + 18), btn_txt, font=FONT_BODY_BOLD, fill=C_WHITE)
        
        draw.text((1120, y + 30), "O que acontece nos bastidores:", font=FONT_SMALL_BOLD, fill=MOD_SECONDARY)
        words = explanation.split()
        line = ""
        ey = y + 60
        for w in words:
            if len(line + " " + w) > 55:
                draw.text((1120, ey), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                ey += 26
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1120, ey), line, font=FONT_BODY, fill=C_TEXT_LIGHT)

    path = os.path.join(TEMP_DIR, "mod8_slide6.png")
    img.save(path)
    return path

def create_slide_7():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • Assistente Guiado", "Passo 4: Assistente de Configuracao Inicial (Setup Wizard)", "Configuracao rapida dos dados fundamentais da loja no primeiro acesso", MOD_ACCENT)
    
    draw_browser_chrome(draw, 80, 260, 1250, 720, "Configuracao do WooCommerce", "https://meusite.com/wp-admin/admin.php?page=wc-admin")
    
    wx, wy = 180, 340
    draw.rounded_rectangle([wx, wy, wx + 1050, wy + 600], radius=12, fill=C_WHITE, outline=(220, 230, 240), width=2)
    
    draw.text((wx + 50, wy + 40), "Bem-vindo ao WooCommerce!", font=FONT_TITLE, fill=(30, 40, 60))
    draw.text((wx + 50, wy + 95), "Vamos configurar as informacoes essenciais para sua loja comecar a vender.", font=FONT_BODY, fill=(100, 110, 130))
    
    steps = ["1. Detalhes da Loja", "2. Industria", "3. Tipos de Produto", "4. Detalhes do Negocio", "5. Tema"]
    for idx, s in enumerate(steps):
        sx = wx + 50 + idx * 200
        sy = wy + 145
        col = MOD_ACCENT if idx == 0 else (200, 210, 220)
        draw.rounded_rectangle([sx, sy, sx + 180, sy + 36], radius=6, fill=col)
        draw.text((sx + 15, sy + 8), s, font=FONT_SMALL_BOLD, fill=C_WHITE if idx == 0 else (80, 90, 100))
        
    fy = wy + 210
    fields = [
        ("Onde sua loja esta sediada?", "Brasil - Sao Paulo"),
        ("Endereco (Linha 1):", "Av. Paulista, 1000 - Sala 42"),
        ("Cidade:", "Sao Paulo"),
        ("CEP / Codigo Postal:", "01310-100"),
        ("Moeda da Loja:", "Real Brasileiro (R$ - BRL)")
    ]
    for lbl, val in fields:
        draw.text((wx + 50, fy), lbl, font=FONT_SMALL_BOLD, fill=(60, 70, 85))
        draw.rounded_rectangle([wx + 50, fy + 24, wx + 650, fy + 60], radius=6, fill=(245, 248, 252), outline=(190, 205, 225), width=1)
        draw.text((wx + 65, fy + 32), val, font=FONT_BODY, fill=(20, 30, 45))
        fy += 70

    draw.rounded_rectangle([wx + 50, fy + 10, wx + 250, fy + 58], radius=6, fill=MOD_ACCENT)
    draw.text((wx + 90, fy + 22), "Continuar ->", font=FONT_BODY_BOLD, fill=C_WHITE)

    # Right explanations
    draw.rounded_rectangle([1370, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1405, 300), "Por que esta etapa e vital?", font=FONT_HEADING, fill=C_WHITE)
    
    points = [
        ("Localizacao Fiscal", "Define a base de calculo de impostos e as zonas de frete prioritarias."),
        ("Moeda Padronizada", "Configurar 'BRL (R$)' garante que os precos serao formatados corretamente (R$ 99,90) e aceitos pelos gateways nacionais."),
        ("CEP de Origem", "O CEP preenchido aqui servira de origem padrao para as cotacoes nos Correios e transportadoras."),
        ("Pode alterar depois?", "Sim! Todas as informacoes podem ser editadas a qualquer momento em WooCommerce > Configuracoes.")
    ]
    dy = 370
    for t, d in points:
        draw_bullet_dot(draw, 1415, dy + 10, radius=5, color=MOD_SECONDARY)
        draw.text((1435, dy), t, font=FONT_BODY_BOLD, fill=MOD_SECONDARY)
        dy += 28
        words = d.split()
        line = ""
        for w in words:
            if len(line + " " + w) > 30:
                draw.text((1435, dy), line, font=FONT_SMALL, fill=C_TEXT_MUTED)
                dy += 22
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1435, dy), line, font=FONT_SMALL, fill=C_TEXT_MUTED)
            dy += 35

    path = os.path.join(TEMP_DIR, "mod8_slide7.png")
    img.save(path)
    return path

def create_slide_8():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • Estrutura Criada", "Paginas Automaticas do WooCommerce", "As 4 paginas criadas automaticamente na instalacao que formam a loja", MOD_ACCENT)
    
    pages = [
        ("LOJA (SHOP)", "/loja", "A vitrine principal onde todos os produtos cadastrados sao exibidos em grade com paginacao, filtros e ordenacao.", (34, 113, 177)),
        ("CARRINHO (CART)", "/carrinho", "Onde o comprador visualiza os itens selecionados, ajusta quantidades, insere cupom de desconto e calcula o frete preliminar.", MOD_ACCENT),
        ("FINALIZAR COMPRA (CHECKOUT)", "/finalizar-compra", "A pagina mais critica: coleta endereco de entrega, dados de faturamento, escolha de envio e processamento do pagamento.", (239, 68, 68)),
        ("MINHA CONTA (MY ACCOUNT)", "/minha-conta", "Area do cliente logado: historico de pedidos anteriores, download de notas fiscais/arquivos e alteracao de senha/endereco.", (16, 185, 129))
    ]
    
    for i, (p_title, p_slug, p_desc, col) in enumerate(pages):
        x = 80 + (i % 2) * 890
        y = 260 + (i // 2) * 350
        
        draw.rounded_rectangle([x, y, x + 850, y + 310], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
        draw.rectangle([x, y, x + 12, y + 310], fill=col)
        
        draw.rounded_rectangle([x + 35, y + 30, x + 115, y + 110], radius=12, fill=col)
        draw.text((x + 55, y + 45), str(i+1), font=FONT_TITLE, fill=C_WHITE)
        
        draw.text((x + 135, y + 35), p_title, font=FONT_HEADING, fill=C_WHITE)
        draw.text((x + 135, y + 75), "URL padrao: " + p_slug, font=FONT_CODE, fill=col)
        
        words = p_desc.split()
        line = ""
        dy = y + 135
        for w in words:
            if len(line + " " + w) > 52:
                draw.text((x + 35, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                dy += 28
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((x + 35, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
            
        draw.rounded_rectangle([x + 35, y + 240, x + 815, y + 285], radius=6, fill=(15, 23, 42))
        draw.text((x + 50, y + 252), "[ATENCAO] Nunca exclua estas paginas! Elas utilizam shortcodes vitais do WooCommerce.", font=FONT_SMALL, fill=(245, 158, 11))

    path = os.path.join(TEMP_DIR, "mod8_slide8.png")
    img.save(path)
    return path

def create_slide_9():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • Visao do Painel", "Novos Menus no Painel WordPress", "A separacao funcional entre operacoes de venda e catalogo de produtos", MOD_ACCENT)
    
    draw.rounded_rectangle([80, 260, 920, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.rounded_rectangle([115, 290, 885, 360], radius=12, fill=MOD_ACCENT)
    draw.text((140, 308), "Menu: WooCommerce (Operacoes da Loja)", font=FONT_HEADING, fill=C_WHITE)
    
    woo_sub = [
        ("Pedidos", "Gerenciar compras recebidas, alterar status para 'Concluido' ou 'Processando'."),
        ("Clientes", "Lista de todos os compradores cadastrados com dados de contato e historico."),
        ("Relatorios / Analises", "Graficos de faturamento bruto, vendas liquidas, itens mais vendidos e devolucoes."),
        ("Cupons", "Criacao de descontos percentuais, fixos ou de frete gratis com regras de uso."),
        ("Configuracoes", "Painel central de regras: enderecos, moedas, metodos de pagamento e frete."),
        ("Status", "Diagnostico tecnico do servidor, versoes do banco e diretorios de log.")
    ]
    for idx, (st, sd) in enumerate(woo_sub):
        sy = 390 + idx * 95
        draw.rounded_rectangle([115, sy, 885, sy + 80], radius=8, fill=(20, 28, 42), outline=(40, 52, 70), width=1)
        draw.text((135, sy + 12), "[Loja] " + st, font=FONT_BODY_BOLD, fill=(180, 205, 255))
        draw.text((135, sy + 44), sd, font=FONT_SMALL, fill=C_TEXT_MUTED)

    draw.rounded_rectangle([960, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.rounded_rectangle([995, 290, 1805, 360], radius=12, fill=(34, 113, 177))
    draw.text((1020, 308), "Menu: Produtos (Gestao do Catalogo)", font=FONT_HEADING, fill=C_WHITE)
    
    prod_sub = [
        ("Todos os Produtos", "Visualizacao tabular do catalogo com foto, SKU, estoque, preco e categoria."),
        ("Adicionar Novo", "Tela completa de cadastro de produtos simples, variaveis, virtuais ou para download."),
        ("Categorias", "Arvore taxonomica para organizacao dos itens (ex: Masculino > Camisas)."),
        ("Tags", "Etiquetas livres para filtros rapidos na busca (ex: algodao, verao, promocao)."),
        ("Atributos", "Configuracao de variacoes globais: Tamanho (P, M, G), Cor, Voltagem (110v, 220v)."),
        ("Comentarios / Avaliacoes", "Moderacao de resenhas e notas deixadas pelos clientes nas paginas de produtos.")
    ]
    for idx, (st, sd) in enumerate(prod_sub):
        sy = 390 + idx * 95
        draw.rounded_rectangle([995, sy, 1805, sy + 80], radius=8, fill=(20, 28, 42), outline=(40, 52, 70), width=1)
        draw.text((1015, sy + 12), "[Item] " + st, font=FONT_BODY_BOLD, fill=(160, 230, 255))
        draw.text((1015, sy + 44), sd, font=FONT_SMALL, fill=C_TEXT_MUTED)

    path = os.path.join(TEMP_DIR, "mod8_slide9.png")
    img.save(path)
    return path

def create_slide_10():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • Configuracoes Essenciais", "Configuracoes Gerais da Loja", "WooCommerce > Configuracoes > Aba Geral: regras de comercio e moeda", MOD_ACCENT)
    
    settings = [
        ("1. Endereco da Loja", [
            ("Linha de Endereco:", "Rua das Lojas, 500"),
            ("Cidade e Estado:", "Sao Paulo - SP"),
            ("Pais:", "Brasil"),
            ("CEP de Origem:", "01000-000")
        ], "Fundamental: este endereco e usado para calcular impostos e calcular o frete nos Correios."),
        
        ("2. Locais de Venda e Entrega", [
            ("Locais de Venda:", "Vender para paises especificos (Brasil)"),
            ("Locais de Entrega:", "Entregar apenas para o Brasil"),
            ("Localizacao Padrao:", "Geolocalizacao ou Endereco da Loja"),
            ("Ativar Impostos:", "Habilitar calculo de taxas se aplicavel")
        ], "Restringir vendas ao Brasil evita pedidos internacionais acidentais sem calculo de frete."),
        
        ("3. Opcoes de Moeda (BRL)", [
            ("Moeda:", "Real Brasileiro (R$)"),
            ("Posicao da Moeda:", "A esquerda com espaco (R$ 99,00)"),
            ("Separador de Milhar:", "Ponto (.)  -> Ex: 1.500,00"),
            ("Separador Decimal:", "Virgula (,) -> Ex: 99,90"),
            ("Numero de Decimais:", "2 decimais")
        ], "Padrao brasileiro obrigatorio para evitar conflitos de formatacao no checkout.")
    ]
    
    for i, (card_title, fields, tip) in enumerate(settings):
        x = 80 + i * 595
        y = 260
        draw.rounded_rectangle([x, y, x + 565, y + 720], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
        
        draw.rounded_rectangle([x + 20, y + 20, x + 545, y + 80], radius=10, fill=(24, 32, 47), outline=MOD_ACCENT, width=1)
        draw.text((x + 35, y + 36), card_title, font=FONT_HEADING, fill=C_WHITE)
        
        fy = y + 110
        for k, v in fields:
            draw.text((x + 30, fy), k, font=FONT_SMALL_BOLD, fill=MOD_SECONDARY)
            draw.rounded_rectangle([x + 30, fy + 24, x + 535, fy + 65], radius=6, fill=(15, 23, 42), outline=(45, 55, 75), width=1)
            draw.text((x + 45, fy + 34), v, font=FONT_BODY, fill=C_TEXT_LIGHT)
            fy += 80
            
        draw.rounded_rectangle([x + 25, y + 570, x + 540, y + 695], radius=10, fill=(20, 30, 42))
        draw.text((x + 40, y + 585), "[IMPORTANCIA]:", font=FONT_SMALL_BOLD, fill=(16, 185, 129))
        words = tip.split()
        line = ""
        ty = y + 615
        for w in words:
            if len(line + " " + w) > 34:
                draw.text((x + 40, ty), line, font=FONT_SMALL, fill=C_TEXT_MUTED)
                ty += 22
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((x + 40, ty), line, font=FONT_SMALL, fill=C_TEXT_MUTED)

    path = os.path.join(TEMP_DIR, "mod8_slide10.png")
    img.save(path)
    return path

def create_slide_11():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • Unidades de Medida", "Configuracoes de Produtos e Medidas", "Aba Produtos: definindo unidades de peso e dimensoes para o frete", MOD_ACCENT)
    
    draw.rounded_rectangle([80, 260, 920, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((115, 295), "Unidades de Medida (Padrao Brasil)", font=FONT_HEADING, fill=C_WHITE)
    
    draw.rounded_rectangle([115, 360, 885, 500], radius=12, fill=(20, 28, 42), outline=(16, 185, 129), width=2)
    draw.text((145, 380), "Unidade de Peso: Quilograma (kg)", font=FONT_HEADING, fill=(16, 185, 129))
    draw.text((145, 425), "Opcoes: kg, g, lbs, oz.", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((145, 455), "-> Escolha SEMPRE 'kg'. Os Correios calculam frete baseado em quilogramas (ex: 0.350 kg).", font=FONT_SMALL, fill=C_TEXT_MUTED)
    
    draw.rounded_rectangle([115, 530, 885, 670], radius=12, fill=(20, 28, 42), outline=MOD_ACCENT, width=2)
    draw.text((145, 550), "Unidade de Dimensoes: Centimetros (cm)", font=FONT_HEADING, fill=MOD_SECONDARY)
    draw.text((145, 595), "Opcoes: m, cm, mm, in, yd.", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((145, 625), "-> Escolha SEMPRE 'cm'. Comprimento, Largura e Altura das caixas sao medidos em centimetros.", font=FONT_SMALL, fill=C_TEXT_MUTED)

    draw.rounded_rectangle([115, 700, 885, 940], radius=12, fill=(20, 28, 42), outline=(45, 55, 75), width=1)
    draw.text((145, 725), "Comportamento ao Adicionar ao Carrinho", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw.text((145, 765), "[X] Redirecionar para o carrinho apos adicao bem-sucedida", font=FONT_SMALL, fill=C_TEXT_LIGHT)
    draw.text((145, 795), "[X] Ativar botoes AJAX de adicao ao carrinho nos arquivos", font=FONT_SMALL, fill=C_TEXT_LIGHT)
    draw.text((145, 840), "Imagem Marcador de Posicao (Placeholder):", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw.text((145, 875), "Imagem padrao que aparece quando o produto nao tem foto.", font=FONT_SMALL, fill=C_TEXT_MUTED)

    draw.rounded_rectangle([960, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((995, 295), "Por que a configuracao correta e CRITICA?", font=FONT_HEADING, fill=C_WHITE)
    
    reasons = [
        ("Calculo de Frete dos Correios / Transportadoras", "Se voce cadastrar um produto como '350' achando que esta em gramas, mas a unidade estiver em 'kg', o sistema tentara cotar 350 QUILOS! O frete resultara em milhares de reais ou erro."),
        ("Calculo de Cubagem Automatico", "O volume da caixa (C x L x A) determina se a transportadora cobrara por peso real ou peso cubado."),
        ("Compatibilidade com Plugins de Envio", "Melhor Envio, Frenet e Correios exigem 'kg' e 'cm' para responder cotacoes em tempo real."),
        ("Regra de Ouro:", "Configure uma vez aqui no inicio para que todos os futuros cadastros de produtos sigam a mesma unidade padronizada.")
    ]
    dy = 370
    for t, d in reasons:
        draw_bullet_dot(draw, 1005, dy + 10, radius=5, color=(245, 158, 11) if "Ouro" in t else MOD_SECONDARY)
        draw.text((1025, dy), t, font=FONT_BODY_BOLD, fill=(245, 158, 11) if "Ouro" in t else MOD_SECONDARY)
        dy += 30
        words = d.split()
        line = ""
        for w in words:
            if len(line + " " + w) > 55:
                draw.text((1025, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                dy += 28
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1025, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
            dy += 45

    path = os.path.join(TEMP_DIR, "mod8_slide11.png")
    img.save(path)
    return path

def create_slide_12():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • E-mails e Notificacoes", "Comunicacao e E-mails Transacionais", "WooCommerce > Configuracoes > Aba E-mails: fluxo de notificacoes automaticas", MOD_ACCENT)
    
    emails = [
        ("Novo Pedido", "Enviado para o Administrador da loja quando um cliente finaliza uma compra.", "Lojista"),
        ("Pedido Cancelado", "Avisa o lojista que um pedido foi cancelado por falta de pagamento ou timeout.", "Lojista"),
        ("Pedido com Falha", "Notifica o lojista quando houve erro no gateway de pagamento.", "Lojista"),
        ("Pedido em Espera", "Enviado ao cliente com instrucoes para pagamento (ex: codigo Pix / Boleto).", "Cliente"),
        ("Processando Pedido", "Enviado ao cliente confirmando que o pagamento foi APROVADO com sucesso.", "Cliente"),
        ("Pedido Concluido", "Enviado ao cliente informando que o pacote foi postado e fornecendo rastreio.", "Cliente"),
        ("Pedido Reembolsado", "Notificacao de estorno de pagamento efetuado.", "Cliente"),
        ("Fatura do Cliente", "Envia por e-mail a relacao detalhada do pedido e link para pagamento.", "Cliente")
    ]
    
    for i, (e_title, e_desc, recipient) in enumerate(emails):
        x = 80 + (i % 2) * 890
        y = 260 + (i // 2) * 175
        
        draw.rounded_rectangle([x, y, x + 850, y + 150], radius=14, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
        
        rec_col = (16, 185, 129) if recipient == "Cliente" else (34, 113, 177)
        draw.rounded_rectangle([x + 25, y + 25, x + 150, y + 62], radius=6, fill=rec_col)
        draw.text((x + 40, y + 33), recipient.upper(), font=FONT_SMALL_BOLD, fill=C_WHITE)
        
        draw.text((x + 170, y + 30), e_title, font=FONT_HEADING, fill=C_WHITE)
        
        words = e_desc.split()
        line = ""
        dy = y + 78
        for w in words:
            if len(line + " " + w) > 58:
                draw.text((x + 25, dy), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                dy += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((x + 25, dy), line, font=FONT_BODY, fill=C_TEXT_MUTED)

    path = os.path.join(TEMP_DIR, "mod8_slide12.png")
    img.save(path)
    return path

def create_slide_13():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • Diagnostico Tecnico", "Verificando o Status do Sistema", "WooCommerce > Status: painel de saude para prevenir erros antes de vender", MOD_ACCENT)
    
    draw.rounded_rectangle([80, 260, 920, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((115, 295), "Itens de Servidor e Ambiente", font=FONT_HEADING, fill=C_WHITE)
    
    server_items = [
        ("Versao do WordPress", "6.x ou superior", True),
        ("Versao do WooCommerce", "Versao mais recente estavel", True),
        ("Versao do PHP", "8.1 / 8.2 (64 bit)", True),
        ("Limite de Memoria PHP", "Minimo 256MB / 512MB", True),
        ("MySQL Versao", "5.7+ ou MariaDB 10.4+", True),
        ("Extensao cURL", "Ativada (comunicacao com APIs)", True),
        ("Extensao OpenSSL", "Ativada (seguranca SSL)", True),
        ("Suporte a Multibyte (mbstring)", "Ativado", True)
    ]
    for idx, (lbl, val, ok) in enumerate(server_items):
        iy = 360 + idx * 72
        draw.rounded_rectangle([115, iy, 885, iy + 60], radius=8, fill=(20, 28, 42), outline=(45, 55, 75), width=1)
        draw_check_icon(draw, 145, iy + 30, radius=12, bg=(16, 185, 129))
        draw.text((175, iy + 18), lbl, font=FONT_BODY_BOLD, fill=C_WHITE)
        draw.text((580, iy + 18), val, font=FONT_SMALL, fill=(16, 185, 129))

    draw.rounded_rectangle([960, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((995, 295), "Paginas, Banco de Dados e Seguranca", font=FONT_HEADING, fill=C_WHITE)
    
    security_items = [
        ("Conexao Segura (HTTPS / SSL)", "Ativo e Valido no Dominio", True),
        ("Tabelas do WooCommerce no MySQL", "16 tabelas wc_ criadas com sucesso", True),
        ("Pagina da Loja configurada", "Pagina com ID vinculado", True),
        ("Pagina do Carrinho configurada", "Shortcode [woocommerce_cart] ok", True),
        ("Pagina de Checkout configurada", "Shortcode [woocommerce_checkout] ok", True),
        ("Pagina Minha Conta configurada", "Shortcode [woocommerce_my_account] ok", True),
        ("Diretorio de Uploads (/wp-content)", "Permissao de escrita (755 / 644)", True),
        ("Registro de Logs (/wp-content/uploads/wc-logs/)", "Gravavel para depuracao de erros", True)
    ]
    for idx, (lbl, val, ok) in enumerate(security_items):
        iy = 360 + idx * 72
        draw.rounded_rectangle([995, iy, 1805, iy + 60], radius=8, fill=(20, 28, 42), outline=(45, 55, 75), width=1)
        draw_check_icon(draw, 1025, iy + 30, radius=12, bg=(16, 185, 129))
        draw.text((1055, iy + 18), lbl, font=FONT_BODY_BOLD, fill=C_WHITE)
        draw.text((1490, iy + 18), val, font=FONT_SMALL, fill=(16, 185, 129))

    path = os.path.join(TEMP_DIR, "mod8_slide13.png")
    img.save(path)
    return path

def create_slide_14():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • Melhores Praticas", "Boas Praticas de Seguranca e Desempenho", "Garantindo uma loja rapida, confiavel e protegida contra ataques", MOD_ACCENT)
    
    practices = [
        ("Cache Inteligente para E-commerce", "Plugins de cache (como WP Super Cache ou LiteSpeed Cache) DEVEM excluir as paginas dinamicas: /carrinho, /finalizar-compra e /minha-conta do cache para nao misturar sessoes de clientes.", (16, 185, 129)),
        ("Backups Diarios Automatizados", "Uma loja virtual nao pode perder pedidos. Configure backups diarios automaticos do banco de dados MySQL e da pasta /wp-content usando UpdraftPlus ou o painel da hospedagem.", (34, 113, 177)),
        ("Autenticacao em Duas Etapas (2FA)", "Proteja as contas de administrador contra invasoes com autenticacao de dois fatores no login do WordPress para garantir a seguranca dos dados financeiros e de clientes.", (239, 68, 68)),
        ("Manter Plugins e Temas Atualizados", "Atualizacoes do WooCommerce trazem correcoes criticas de seguranca e compatibilidade com navegadores e APIs de pagamento.", MOD_ACCENT)
    ]
    
    for i, (title, desc, col) in enumerate(practices):
        y = 260 + i * 175
        draw.rounded_rectangle([80, y, 1840, y + 150], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
        draw.rectangle([80, y, 92, y + 150], fill=col)
        
        draw.ellipse([120, y + 35, 180, y + 95], fill=col)
        draw.text((142, y + 48), str(i+1), font=FONT_HEADING, fill=C_WHITE)
        
        draw.text((210, y + 35), title, font=FONT_HEADING, fill=C_WHITE)
        
        words = desc.split()
        line = ""
        dy = y + 78
        for w in words:
            if len(line + " " + w) > 95:
                draw.text((210, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                dy += 26
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((210, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)

    path = os.path.join(TEMP_DIR, "mod8_slide14.png")
    img.save(path)
    return path

def create_slide_15():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 8 • Checklist de Conclusao", "Resumo do Modulo e Proximos Passos", "Tudo o que foi configurado com sucesso e o que faremos no proximo modulo", MOD_ACCENT)
    
    draw.rounded_rectangle([80, 260, 1050, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((115, 295), "Checklist de Instalacao do WooCommerce", font=FONT_HEADING, fill=C_WHITE)
    
    checklist = [
        "Plugin WooCommerce oficial da Automattic instalado e ativo",
        "Assistente inicial concluido com pais 'Brasil' e moeda 'R$ (BRL)'",
        "Paginas essenciais criadas: Loja, Carrinho, Checkout e Minha Conta",
        "Unidades de medida padronizadas: Peso em 'kg' e Dimensoes em 'cm'",
        "Menus 'WooCommerce' e 'Produtos' disponiveis e acessiveis no painel",
        "Status do Sistema verificado: sem alertas vermelhos de PHP ou MySQL",
        "Politicas de retencao de dados e criacao de conta no checkout ativas"
    ]
    for idx, item in enumerate(checklist):
        iy = 360 + idx * 80
        draw.rounded_rectangle([115, iy, 1015, iy + 65], radius=10, fill=(20, 30, 42), outline=(16, 185, 129), width=1)
        draw_check_icon(draw, 145, iy + 32, radius=12, bg=(16, 185, 129))
        draw.text((175, iy + 20), item, font=FONT_BODY, fill=C_TEXT_LIGHT)

    draw.rounded_rectangle([1090, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=MOD_ACCENT, width=2)
    draw.rounded_rectangle([1125, 290, 1805, 360], radius=10, fill=MOD_ACCENT)
    draw.text((1150, 310), "PROXIMO PASSO: MODULO 09", font=FONT_HEADING, fill=C_WHITE)
    
    draw.text((1125, 390), "Plugin Mercado Livre & Mercado Pago", font=FONT_TITLE, fill=C_WHITE)
    draw.text((1125, 455), "Agora que a base do e-commerce esta pronta, vamos integrar o motor de vendas e pagamentos:", font=FONT_BODY, fill=C_TEXT_MUTED)
    
    next_topics = [
        ("Mercado Pago", "Ativar recebimentos instantaneos por Pix com QR Code, cartao parcelado e boleto."),
        ("Mercado Livre", "Sincronizar produtos, estoque e vendas entre a loja virtual e o maior marketplace."),
        ("Checkout Transparente", "Comprador paga sem sair do seu site, aumentando drasticamente a taxa de conversao.")
    ]
    ny = 525
    for nt, nd in next_topics:
        draw.rounded_rectangle([1125, ny, 1805, ny + 115], radius=12, fill=(24, 32, 47), outline=(45, 55, 75), width=1)
        draw.text((1150, ny + 18), "->  " + nt, font=FONT_HEADING, fill=(255, 214, 0))
        draw.text((1150, ny + 60), nd, font=FONT_BODY, fill=C_TEXT_LIGHT)
        ny += 135

    path = os.path.join(TEMP_DIR, "mod8_slide15.png")
    img.save(path)
    return path

def build_presentation_8():
    print("Generating slides for 8 instalacao de woocomerce.pptx...")
    slide_funcs = [
        create_slide_1, create_slide_2, create_slide_3, create_slide_4,
        create_slide_5, create_slide_6, create_slide_7, create_slide_8,
        create_slide_9, create_slide_10, create_slide_11, create_slide_12,
        create_slide_13, create_slide_14, create_slide_15
    ]
    
    prs = Presentation("6 wordpress.pptx")
    rIds = [prs.slides._sldIdLst[i].rId for i in range(len(prs.slides))]
    for rId in rIds:
        prs.part.drop_rel(rId)
    prs.slides._sldIdLst.clear()
    
    blank_layout = prs.slide_layouts[6]
    
    notes_scripts = [
        "Apresentacao do Modulo 8: Instalacao e Configuracao do WooCommerce. Explicar aos alunos o objetivo da aula: transformar o WordPress instalado no Modulo 6 em uma loja virtual completa.",
        "Explicar a arquitetura em camadas do WooCommerce e por que ele e lider absoluto no mercado global de e-commerce, destacando a posse real dos dados e a ausencia de comissoes por venda.",
        "Destacar a importancia do PHP 8.x, memoria de 256MB+ e do certificado SSL HTTPS ativo. Sem SSL nenhum gateway de pagamento aprova transacoes.",
        "Passo a passo no painel: mostrar onde fica o menu lateral de Plugins e o botao 'Adicionar Novo' que conecta ao repositorio oficial.",
        "Demonstrar a pesquisa por 'WooCommerce'. Alertar os alunos para checarem se o autor e a Automattic e se ha mais de 5 milhoes de instalacoes ativas.",
        "Explicar os 3 estados do botao: Instalar Agora, Instalando (download e descompactacao em /wp-content/plugins/) e Ativar (criacao das tabelas no MySQL).",
        "Guiar os alunos pelo Assistente Inicial (Setup Wizard). Mostrar a importancia de definir Brasil e a moeda Real (R$ BRL) logo de inicio.",
        "Mostrar as 4 paginas essenciais criadas automaticamente (Loja, Carrinho, Finalizar Compra e Minha Conta) e alertar severamente para jamais exclui-las.",
        "Explicar a separacao logica no painel do WordPress entre o menu 'WooCommerce' (operacoes, pedidos e clientes) e o menu 'Produtos' (catalogo).",
        "Passar pelas Configuracoes Gerais: endereco fisico da loja para frete, restricao de venda apenas para o Brasil e formatacao correta da moeda BRL.",
        "Atencao total nesta aula para as Unidades de Medida: Peso SEMPRE em 'kg' e Dimensoes SEMPRE em 'cm'. Explicar o erro catastrofico que ocorre se cadastrar gramas como kg.",
        "Explicar o fluxo de e-mails transacionais automaticos enviados pelo WooCommerce para cliente e lojista em cada etapa do pedido.",
        "Mostrar aos alunos como acessar 'WooCommerce > Status' para auditar a saude do servidor e confirmar que todos os itens estao verdes.",
        "Abordar seguranca e performance: regras de cache para e-commerce (nunca cachear carrinho/checkout), backups diarios e autenticacao em dois fatores.",
        "Revisao final com o checklist de conclusao do Modulo 8 e introducao do Modulo 9: Integracao com Mercado Livre e Mercado Pago."
    ]
    
    for idx, func in enumerate(slide_funcs):
        img_path = func()
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(img_path, 0, 0, prs.slide_width, prs.slide_height)
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = notes_scripts[idx]
        print(f"  Slide {idx+1}/{len(slide_funcs)} generated.")

    out_file = "8 instalacao de woocomerce.pptx"
    prs.save(out_file)
    print(f"Presentation saved successfully as '{out_file}' ({os.path.getsize(out_file)} bytes).")

if __name__ == "__main__":
    build_presentation_8()
