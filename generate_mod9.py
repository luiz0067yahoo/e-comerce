"""
Generator for: 9 plugin mercado livre.pptx
"""

import os
from PIL import Image, ImageDraw, ImageFont
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt

from slide_builder import (
    TEMP_DIR, FONT_HERO, FONT_TITLE, FONT_SUBTITLE, FONT_HEADING,
    FONT_BODY, FONT_BODY_BOLD, FONT_SMALL, FONT_SMALL_BOLD, FONT_CODE,
    C_DARK_BG, C_CARD_BG, C_CARD_BORDER, C_TEXT_LIGHT, C_TEXT_MUTED, C_WHITE,
    create_base_canvas, draw_header, draw_browser_chrome,
    draw_wp_sidebar, draw_arrow_pointer, draw_check_icon, draw_cross_icon,
    draw_bullet_dot, draw_star_badge
)

MELI_YELLOW = (255, 214, 0)
MP_BLUE = (0, 158, 227)
MOD_ACCENT = MP_BLUE

def create_slide_1():
    img, draw = create_base_canvas()
    draw.rounded_rectangle([120, 160, 520, 215], radius=12, fill=MELI_YELLOW)
    draw.text((140, 172), "MODULO 09 • INTEGRACAO", font=FONT_HEADING, fill=(20, 20, 20))
    
    draw.text((120, 250), "Plugin Mercado Livre", font=FONT_HERO, fill=C_WHITE)
    draw.text((120, 335), "& Gateway Mercado Pago", font=FONT_HERO, fill=MELI_YELLOW)
    
    draw.text((120, 440), "Vendas Multicanal no Maior Marketplace da America Latina e Pagamentos Transparentes", font=FONT_SUBTITLE, fill=C_TEXT_MUTED)
    
    features = [
        ("Checkout Transparente", "Receba via Pix com QR Code dinamico, Cartao de Credito parcelado e Boleto sem redirecionar o cliente."),
        ("Sincronizacao de Catalogo", "Publique produtos do WooCommerce diretamente no Mercado Livre com 1 clique."),
        ("Estoque Unificado", "Uma venda no Mercado Livre baixa automaticamente o estoque na sua loja virtual, evitando furos.")
    ]
    for i, (title, desc) in enumerate(features):
        x = 120 + i * 570
        y = 540
        draw.rounded_rectangle([x, y, x + 530, y + 260], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
        draw.ellipse([x + 30, y + 30, x + 80, y + 80], fill=MP_BLUE)
        draw.text((x + 46, y + 38), str(i+1), font=FONT_HEADING, fill=C_WHITE)
        draw.text((x + 100, y + 42), title, font=FONT_HEADING, fill=C_TEXT_LIGHT)
        
        words = desc.split()
        line = ""
        dy = y + 105
        for w in words:
            if len(line + " " + w) > 38:
                draw.text((x + 30, dy), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                dy += 28
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((x + 30, dy), line, font=FONT_BODY, fill=C_TEXT_MUTED)

    draw.rounded_rectangle([120, 850, 1800, 930], radius=12, fill=(24, 32, 47), outline=MP_BLUE, width=1)
    draw_check_icon(draw, 150, 890, radius=14, bg=(16, 185, 129))
    draw.text((180, 875), "Objetivo: Conectar o WooCommerce as APIs oficiais do Mercado Livre e Mercado Pago", font=FONT_BODY_BOLD, fill=(16, 185, 129))
    draw.text((1300, 875), "Curso Pratico de E-commerce", font=FONT_BODY, fill=C_TEXT_MUTED)

    path = os.path.join(TEMP_DIR, "mod9_slide1.png")
    img.save(path)
    return path

def create_slide_2():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Estrategia Comercial", "Visao Geral da Integracao Omnichannel", "Como combinar a forca do marketplace com a liberdade da loja propria", MP_BLUE)
    
    # 3 big pillars
    pillars = [
        ("1. SUA LOJA VIRTUAL", "WooCommerce", [
            "Controle total da marca e identidade visual",
            "Base de clientes e e-mails 100% sua",
            "Sem comissao por venda sobre os produtos",
            "Espaco para fidelizacao e recompra direta"
        ], (127, 84, 179)),
        
        ("2. GATEWAY DE PAGAMENTO", "Mercado Pago", [
            "Checkout transparente sem sair da loja",
            "Pix com aprovacao instantanea em segundos",
            "Cartao em ate 12x com protecao antifraude",
            "Boleto bancario com baixa automatica"
        ], MP_BLUE),
        
        ("3. CANAL DE TRACAO", "Mercado Livre", [
            "Milhoes de compradores navegando diariamente",
            "Tracao imediata para produtos recem-lancados",
            "Logistica e confianca do ecossistema Meli",
            "Sincronizacao de estoque bidirecional"
        ], MELI_YELLOW)
    ]
    
    for i, (title, sub, bullets, col) in enumerate(pillars):
        x = 80 + i * 595
        y = 260
        draw.rounded_rectangle([x, y, x + 565, y + 720], radius=16, fill=C_CARD_BG, outline=col, width=2)
        
        draw.rounded_rectangle([x + 20, y + 20, x + 545, y + 100], radius=12, fill=(20, 28, 42), outline=col, width=1)
        draw.text((x + 35, y + 32), title, font=FONT_BODY_BOLD, fill=C_WHITE)
        draw.text((x + 35, y + 62), sub, font=FONT_TITLE, fill=col)
        
        by = y + 140
        for b in bullets:
            draw_check_icon(draw, x + 45, by + 12, radius=10, bg=col if col != MELI_YELLOW else (200, 160, 0))
            words = b.split()
            line = ""
            for w in words:
                if len(line + " " + w) > 34:
                    draw.text((x + 68, by), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                    by += 26
                    line = w
                else:
                    line = (line + " " + w).strip()
            if line:
                draw.text((x + 68, by), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                by += 38

    path = os.path.join(TEMP_DIR, "mod9_slide2.png")
    img.save(path)
    return path

def create_slide_3():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Arquitetura", "Os Dois Pilares da Integracao", "Compreendendo a funcao exata de cada componente do ecossistema", MP_BLUE)
    
    # Diagram: Loja no meio, Mercado Pago na esquerda, Mercado Livre na direita
    # Left box: Mercado Pago
    draw.rounded_rectangle([80, 270, 580, 950], radius=16, fill=C_CARD_BG, outline=MP_BLUE, width=2)
    draw.rounded_rectangle([105, 295, 555, 365], radius=10, fill=MP_BLUE)
    draw.text((130, 315), "PILAR A: MERCADO PAGO", font=FONT_HEADING, fill=C_WHITE)
    draw.text((105, 385), "Funcao: Gateway Financeiro", font=FONT_HEADING, fill=MP_BLUE)
    
    mp_details = [
        "Processa as transacoes na sua loja propria.",
        "Gera o QR Code do Pix dinamicamente.",
        "Valida o cartao de credito com sistema antifraude.",
        "Emite webhook que avisa o WooCommerce quando o cliente pagou, mudando o status para 'Processando'."
    ]
    dy = 445
    for item in mp_details:
        draw_bullet_dot(draw, 120, dy + 10, radius=5, color=MP_BLUE)
        words = item.split()
        line = ""
        for w in words:
            if len(line + " " + w) > 32:
                draw.text((140, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                dy += 26
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((140, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
            dy += 45

    # Center box: WooCommerce Hub
    draw.rounded_rectangle([620, 270, 1300, 950], radius=16, fill=(24, 32, 47), outline=(127, 84, 179), width=2)
    draw.rounded_rectangle([645, 295, 1275, 365], radius=10, fill=(127, 84, 179))
    draw.text((670, 315), "CENTRO DE COMANDO: WOOCOMMERCE", font=FONT_HEADING, fill=C_WHITE)
    
    draw.text((645, 385), "Gestao Centralizada", font=FONT_TITLE, fill=C_WHITE)
    draw.text((645, 440), "Todos os produtos, precos, fotos e pedidos ficam centralizados no seu WordPress.", font=FONT_BODY, fill=C_TEXT_MUTED)
    
    draw.rounded_rectangle([645, 510, 1275, 680], radius=12, fill=(15, 23, 42), outline=(51, 65, 85), width=1)
    draw.text((670, 530), "Fluxo de Sincronizacao:", font=FONT_HEADING, fill=(16, 185, 129))
    draw.text((670, 575), "<- 1. Mercado Pago envia confirmacao de pagamento", font=FONT_BODY_BOLD, fill=MP_BLUE)
    draw.text((670, 615), "-> 2. WooCommerce atualiza estoque em tempo real", font=FONT_BODY_BOLD, fill=(127, 84, 179))
    draw.text((670, 655), "<- 3. Mercado Livre envia vendas feitas no marketplace", font=FONT_BODY_BOLD, fill=MELI_YELLOW)

    draw.rounded_rectangle([645, 720, 1275, 920], radius=12, fill=(35, 45, 62))
    draw.text((670, 745), "Vantagem Operacional:", font=FONT_HEADING, fill=C_WHITE)
    draw.text((670, 785), "Sua equipe opera apenas UM painel administrativo,", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((670, 815), "sem precisar gerenciar abas separadas para o marketplace.", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((670, 855), "Economia de tempo e eliminacao de erros humanos.", font=FONT_BODY, fill=(16, 185, 129))

    # Right box: Mercado Livre
    draw.rounded_rectangle([1340, 270, 1840, 950], radius=16, fill=C_CARD_BG, outline=MELI_YELLOW, width=2)
    draw.rounded_rectangle([1365, 295, 1815, 365], radius=10, fill=MELI_YELLOW)
    draw.text((1390, 315), "PILAR B: MERCADO LIVRE", font=FONT_HEADING, fill=(20, 20, 20))
    draw.text((1365, 385), "Funcao: Marketplace de Vendas", font=FONT_HEADING, fill=MELI_YELLOW)
    
    meli_details = [
        "Vitrine com trafego massivo de milhoes de visitas.",
        "Importa seus produtos diretamente do WooCommerce.",
        "Sincroniza precos, fotos e descricoes.",
        "Quando um item e vendido no Mercado Livre, baixa o estoque da sua loja automaticamente."
    ]
    dy = 445
    for item in meli_details:
        draw_bullet_dot(draw, 1380, dy + 10, radius=5, color=MELI_YELLOW)
        words = item.split()
        line = ""
        for w in words:
            if len(line + " " + w) > 28:
                draw.text((1400, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                dy += 26
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1400, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
            dy += 45

    path = os.path.join(TEMP_DIR, "mod9_slide3.png")
    img.save(path)
    return path

def create_slide_4():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Passo a Passo", "Passo 1: Instalar o Plugin Oficial do Mercado Pago", "Localizando e ativando o plugin oficial de pagamentos no WordPress", MP_BLUE)
    
    draw_browser_chrome(draw, 80, 260, 1250, 720, "Adicionar Plugins - Mercado Pago", "https://meusite.com/wp-admin/plugin-install.php?s=Mercado+Pago")
    draw_wp_sidebar(draw, 80, 308, 250, 672, active_item="Plugins")
    
    draw.rectangle([330, 308, 1330, 980], fill=(245, 247, 250))
    draw.text((370, 340), "Adicionar Plugins", font=FONT_TITLE, fill=(30, 40, 55))
    
    draw.rounded_rectangle([800, 340, 1280, 385], radius=6, fill=C_WHITE, outline=(180, 190, 205), width=2)
    draw.text((815, 348), "Pesquisar: Mercado Pago", font=FONT_BODY_BOLD, fill=(40, 50, 70))
    
    # Plugin Card
    card_x, card_y = 370, 430
    draw.rounded_rectangle([card_x, card_y, card_x + 880, card_y + 450], radius=12, fill=C_WHITE, outline=(210, 220, 230), width=2)
    
    draw.rounded_rectangle([card_x + 30, card_y + 30, card_x + 120, card_y + 120], radius=16, fill=MP_BLUE)
    draw.text((card_x + 42, card_y + 45), "MP", font=FONT_HERO, fill=C_WHITE)
    
    draw.text((card_x + 150, card_y + 35), "Mercado Pago payments for WooCommerce", font=FONT_HEADING, fill=(20, 30, 45))
    draw.text((card_x + 150, card_y + 75), "Por Mercado Pago (Autor Oficial Verificado)", font=FONT_BODY_BOLD, fill=MP_BLUE)
    
    desc = "Aceite pagamentos com Cartao de Credito em ate 12x, Pix com aprovacao instantanea e Boleto Bancario com o Checkout Transparente oficial do Mercado Pago."
    draw.text((card_x + 30, card_y + 145), desc[:75], font=FONT_BODY, fill=(70, 80, 95))
    draw.text((card_x + 30, card_y + 180), desc[75:], font=FONT_BODY, fill=(70, 80, 95))
    
    draw_star_badge(draw, card_x + 30, card_y + 250, count=5)
    draw.text((card_x + 150, card_y + 248), "(Mais de 100.000 instalacoes ativas)", font=FONT_BODY_BOLD, fill=(234, 150, 20))
    
    draw_check_icon(draw, card_x + 45, card_y + 298, radius=12, bg=(16, 185, 129))
    draw.text((card_x + 65, card_y + 288), "Compativel com a sua versao do WooCommerce", font=FONT_BODY_BOLD, fill=(16, 185, 129))
    
    draw.rounded_rectangle([card_x + 650, card_y + 35, card_x + 840, card_y + 85], radius=6, fill=(34, 113, 177))
    draw.text((card_x + 685, card_y + 48), "Instalar Agora", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw_arrow_pointer(draw, (card_x + 750, card_y + 160), (card_x + 750, card_y + 95), color=(239, 68, 68), width=5)

    # Right column instructions
    draw.rounded_rectangle([1370, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1405, 300), "Destaques do Plugin", font=FONT_HEADING, fill=C_WHITE)
    
    items = [
        ("Desenvolvedor Oficial", "Instale apenas o plugin desenvolvido diretamente pelo 'Mercado Pago'. Nao utilize plugins de terceiros."),
        ("Checkout Transparente", "O cliente nao e redirecionado para fora do seu site. Ele digita os dados e paga diretamente na sua pagina."),
        ("Pix com QR Code", "Cria QR Code em tempo real na tela e chave 'Copia e Cola'."),
        ("Notificacao Instantanea", "O gateway avisa seu site no exato segundo em que o Pix e pago.")
    ]
    dy = 370
    for t, d in items:
        draw_bullet_dot(draw, 1415, dy + 10, radius=5, color=MP_BLUE)
        draw.text((1435, dy), t, font=FONT_BODY_BOLD, fill=MP_BLUE)
        dy += 28
        words = d.split()
        line = ""
        for w in words:
            if len(line + " " + w) > 30:
                draw.text((1435, dy), line, font=FONT_SMALL, fill=C_TEXT_MUTED)
                dy += 22
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1435, dy), line, font=FONT_SMALL, fill=C_TEXT_MUTED)
            dy += 35

    path = os.path.join(TEMP_DIR, "mod9_slide4.png")
    img.save(path)
    return path

def create_slide_5():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Portal Dev", "Passo 2: Criando Conta de Desenvolvedor", "Acessando o portal do Mercado Pago Developers para criar sua aplicacao", MP_BLUE)
    
    draw_browser_chrome(draw, 80, 260, 1250, 720, "Mercado Pago Developers", "https://www.mercadopago.com.br/developers/panel/app")
    
    # Dev Portal Mockup
    px, py = 120, 320
    draw.rounded_rectangle([px, py, px + 1170, py + 620], radius=12, fill=(245, 248, 252), outline=(210, 225, 240), width=2)
    
    # Dev portal header
    draw.rectangle([px, py, px + 1170, py + 70], fill=MP_BLUE)
    draw.text((px + 40, py + 20), "Mercado Pago | Developers", font=FONT_HEADING, fill=C_WHITE)
    draw.text((px + 900, py + 25), "Minhas Integracoes", font=FONT_BODY_BOLD, fill=C_WHITE)
    
    # "Criar Aplicacao" Card
    cx, cy = px + 60, py + 120
    draw.rounded_rectangle([cx, cy, cx + 1050, cy + 440], radius=12, fill=C_WHITE, outline=(220, 230, 240), width=2)
    draw.text((cx + 40, cy + 35), "Criar Nova Aplicacao", font=FONT_TITLE, fill=(30, 45, 65))
    draw.text((cx + 40, cy + 85), "Preencha os detalhes para obter suas chaves de API e conectar ao WooCommerce", font=FONT_BODY, fill=(100, 115, 130))
    
    fields = [
        ("Nome da Aplicacao:", "Minha Loja Virtual WooCommerce"),
        ("Tipo de Integracao:", "Checkout Transparente (E-commerce)"),
        ("Plataforma de E-commerce:", "WooCommerce / WordPress"),
        ("Ambiente Inicial:", "Modo Sandbox (Teste) e Producao")
    ]
    fy = cy + 140
    for lbl, val in fields:
        draw.text((cx + 40, fy), lbl, font=FONT_BODY_BOLD, fill=(60, 75, 95))
        draw.rounded_rectangle([cx + 340, fy - 6, cx + 980, fy + 36], radius=6, fill=(240, 245, 250), outline=(190, 205, 225), width=1)
        draw.text((cx + 360, fy + 2), val, font=FONT_BODY, fill=(20, 30, 45))
        fy += 60

    # Submit button
    draw.rounded_rectangle([cx + 340, fy + 15, cx + 580, fy + 65], radius=8, fill=MP_BLUE)
    draw.text((cx + 380, fy + 28), "Criar Aplicacao ->", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw_arrow_pointer(draw, (cx + 660, fy + 40), (cx + 590, fy + 40), color=(239, 68, 68), width=5)

    # Right column instructions
    draw.rounded_rectangle([1370, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1405, 300), "Etapas do Portal", font=FONT_HEADING, fill=C_WHITE)
    
    steps = [
        "1. Acesse o site oficial 'mercadopago.com.br/developers'.",
        "2. Faca login usando o mesmo e-mail e senha da sua conta do Mercado Livre / Mercado Pago.",
        "3. Clique em 'Suas integracoes' e depois no botao 'Criar aplicacao'.",
        "4. No campo 'Plataforma', selecione 'WooCommerce'.",
        "5. Ao concluir, a aplicacao gerara o par de credenciais de teste e producao."
    ]
    dy = 370
    for s in steps:
        words = s.split()
        line = ""
        for w in words:
            if len(line + " " + w) > 30:
                draw.text((1405, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                dy += 28
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1405, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
            dy += 42

    path = os.path.join(TEMP_DIR, "mod9_slide5.png")
    img.save(path)
    return path

def create_slide_6():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Seguranca de API", "Passo 3: Obtendo as Credenciais de API", "Public Key e Access Token: as chaves que autorizam a comunicacao segura", MP_BLUE)
    
    # 2 Big Cards: Public Key and Access Token
    # Card 1: Public Key
    draw.rounded_rectangle([80, 260, 920, 680], radius=16, fill=C_CARD_BG, outline=MP_BLUE, width=2)
    draw.rounded_rectangle([115, 290, 450, 340], radius=8, fill=MP_BLUE)
    draw.text((135, 304), "CHAVE PUBLICA (PUBLIC KEY)", font=FONT_SMALL_BOLD, fill=C_WHITE)
    
    draw.text((115, 370), "APP_USR-782a1b9c-4821-4f11-9a7c-abcdef123456", font=FONT_CODE, fill=(160, 220, 255))
    
    desc_pk = [
        "Identifica sua loja publicamente perante o Mercado Pago.",
        "Pode ser enviada ao navegador do cliente sem risco de seguranca.",
        "E utilizada na pagina de checkout para renderizar os campos seguros de cartao de credito (tokenizacao)."
    ]
    dy = 440
    for p in desc_pk:
        draw_check_icon(draw, 135, dy + 10, radius=10, bg=MP_BLUE)
        words = p.split()
        line = ""
        for w in words:
            if len(line + " " + w) > 50:
                draw.text((160, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                dy += 26
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((160, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
            dy += 40

    # Card 2: Access Token
    draw.rounded_rectangle([960, 260, 1840, 680], radius=16, fill=C_CARD_BG, outline=(239, 68, 68), width=2)
    draw.rounded_rectangle([995, 290, 1370, 340], radius=8, fill=(239, 68, 68))
    draw.text((1015, 304), "TOKEN DE ACESSO (ACCESS TOKEN)", font=FONT_SMALL_BOLD, fill=C_WHITE)
    
    draw.text((995, 370), "APP_USR-9876543210123456-102122-abcdef0123456789-123456789", font=FONT_CODE, fill=(255, 180, 180))
    
    desc_at = [
        "SENHA SECRETA da sua aplicacao. Nunca compartilhe nem exponha no front-end.",
        "Autoriza o servidor do WooCommerce a cobrar clientes, consultar saldo e criar reembolsos.",
        "Fica armazenada exclusivamente no banco de dados do seu WordPress, protegida no backend."
    ]
    dy = 440
    for p in desc_at:
        draw_cross_icon(draw, 1015, dy + 10, radius=10, bg=(239, 68, 68))
        words = p.split()
        line = ""
        for w in words:
            if len(line + " " + w) > 50:
                draw.text((1040, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                dy += 26
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1040, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
            dy += 40

    # Bottom alert box
    draw.rounded_rectangle([80, 720, 1840, 960], radius=16, fill=(24, 32, 47), outline=(245, 158, 11), width=2)
    draw.text((120, 755), "[REGRA DE OURO DA INTEGRACAO]", font=FONT_HEADING, fill=(245, 158, 11))
    draw.text((120, 805), "Existem dois conjuntos de chaves no painel do Mercado Pago:", font=FONT_BODY_BOLD, fill=C_WHITE)
    draw.text((120, 845), "1. Credenciais de Teste (Sandbox): Comecam com 'TEST-' -> Usadas para simular compras sem dinheiro real.", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((120, 885), "2. Credenciais de Producao: Comecam com 'APP_USR-' -> Usadas quando a loja estiver pronta para faturar.", font=FONT_BODY, fill=(16, 185, 129))

    path = os.path.join(TEMP_DIR, "mod9_slide6.png")
    img.save(path)
    return path

def create_slide_7():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Ambientes", "Passo 4: Modo Sandbox (Teste) vs Producao", "Como testar todo o fluxo de pagamento antes de cobrar clientes de verdade", MP_BLUE)
    
    # 2 Comparison Columns
    # Left: Sandbox
    draw.rounded_rectangle([80, 260, 920, 980], radius=16, fill=C_CARD_BG, outline=(245, 158, 11), width=2)
    draw.rounded_rectangle([115, 290, 885, 360], radius=12, fill=(245, 158, 11))
    draw.text((140, 310), "MODO SANDBOX (TESTE / HOMOLOGACAO)", font=FONT_HEADING, fill=(20, 20, 20))
    
    sb_points = [
        ("Finalidade:", "Validar se o plugin esta instalado corretamente, se o webhook de retorno funciona e se os pedidos mudam de status."),
        ("Cartoes Ficticios:", "O Mercado Pago fornece numeros de cartoes de teste (ex: 4024 0071 ... com bandeira Visa) para simular pagamentos aprovados, recusados ou em analise."),
        ("Pix Simulado:", "O QR Code e gerado normalmente e pode ser aprovado clicando em um botao simulador no painel."),
        ("Dinheiro Real:", "Nenhum centavo real e movimentado."),
        ("Quando usar:", "Durante toda a fase de criacao e configuracao da loja virtual.")
    ]
    dy = 390
    for lbl, txt in sb_points:
        draw.text((115, dy), lbl, font=FONT_BODY_BOLD, fill=(245, 158, 11))
        words = txt.split()
        line = ""
        ty = dy + 28
        for w in words:
            if len(line + " " + w) > 50:
                draw.text((115, ty), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                ty += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((115, ty), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
            dy = ty + 40

    # Right: Producao
    draw.rounded_rectangle([960, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=(16, 185, 129), width=2)
    draw.rounded_rectangle([995, 290, 1805, 360], radius=12, fill=(16, 185, 129))
    draw.text((1020, 310), "MODO PRODUCAO (AO VIVO / VENDAS REAIS)", font=FONT_HEADING, fill=C_WHITE)
    
    prod_points = [
        ("Finalidade:", "Receber pagamentos reais dos clientes finais e transferir os valores para sua conta bancaria."),
        ("Cartoes Reais:", "Aceita cartoes de credito emitidos por qualquer banco brasileiro (Nubank, Itau, Bradesco, etc.)."),
        ("Pix Oficial:", "Gera chaves Pix oficiais pelo Banco Central com leitura pelo app de qualquer instituicao financeira."),
        ("Taxas e Prazos:", "Aplica as tarifas padrao do Mercado Pago (ex: Pix ~0.99%, Cartao ~3.99% a 4.99% conforme prazo de liberacao)."),
        ("Ativacao:", "So ative este modo apos testar com sucesso todas as etapas no Sandbox!")
    ]
    dy = 390
    for lbl, txt in prod_points:
        draw.text((995, dy), lbl, font=FONT_BODY_BOLD, fill=(16, 185, 129))
        words = txt.split()
        line = ""
        ty = dy + 28
        for w in words:
            if len(line + " " + w) > 50:
                draw.text((995, ty), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                ty += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((995, ty), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
            dy = ty + 40

    path = os.path.join(TEMP_DIR, "mod9_slide7.png")
    img.save(path)
    return path

def create_slide_8():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Configuracao no WP", "Passo 5: Conectando as Credenciais no WooCommerce", "WooCommerce > Configuracoes > Pagamentos > Mercado Pago", MP_BLUE)
    
    draw_browser_chrome(draw, 80, 260, 1250, 720, "Configuracoes de Pagamento WooCommerce", "https://meusite.com/wp-admin/admin.php?page=wc-settings&tab=checkout&section=woo-mercado-pago-basic")
    
    # Form mockup
    fx, fy = 120, 320
    draw.rounded_rectangle([fx, fy, fx + 1170, py:=640], radius=12, fill=C_WHITE, outline=(210, 220, 230), width=2)
    
    draw.text((fx + 40, fy + 30), "Credenciais do Mercado Pago", font=FONT_TITLE, fill=(30, 45, 65))
    draw.text((fx + 40, fy + 80), "Cole abaixo as credenciais obtidas no Portal de Desenvolvedores:", font=FONT_BODY, fill=(100, 115, 130))
    
    cred_fields = [
        ("Chave Publica de Producao (Public Key):", "APP_USR-782a1b9c-4821-4f11-9a7c-abcdef123456"),
        ("Token de Acesso de Producao (Access Token):", "APP_USR-9876543210123456-102122-abcdef0123456789-123456789"),
        ("Chave Publica de Teste (Sandbox Public Key):", "TEST-3498a12b-7890-41ef-8b12-987654fedcba"),
        ("Token de Acesso de Teste (Sandbox Access Token):", "TEST-1234567890123456-102122-987654321abcdef-123456789")
    ]
    y_f = fy + 130
    for lbl, val in cred_fields:
        draw.text((fx + 40, y_f), lbl, font=FONT_SMALL_BOLD, fill=(60, 75, 95))
        draw.rounded_rectangle([fx + 40, y_f + 24, fx + 1050, y_f + 65], radius=6, fill=(245, 248, 252), outline=(190, 205, 225), width=1)
        draw.text((fx + 55, y_f + 34), val, font=FONT_CODE, fill=(20, 30, 45))
        y_f += 85

    # Save button
    draw.rounded_rectangle([fx + 40, y_f + 20, fx + 260, y_f + 70], radius=6, fill=MP_BLUE)
    draw.text((fx + 70, y_f + 32), "Salvar Alteracoes", font=FONT_BODY_BOLD, fill=C_WHITE)

    # Right column
    draw.rounded_rectangle([1370, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((1405, 300), "Dicas de Validacao", font=FONT_HEADING, fill=C_WHITE)
    
    tips = [
        ("Nao deixe espacos", "Ao copiar do portal, certifique-se de nao copiar espacos em branco no inicio ou final da chave."),
        ("Teste de Conexao", "O plugin valida automaticamente se as credenciais sao autenticas ao clicar em 'Salvar'."),
        ("Modo de Teste Ativo", "Marque a caixa 'Ativar modo de teste' para usar as credenciais TEST-."),
        ("Loja Aprovada", "Apos testar, desmarque o modo de teste para iniciar o processamento em producao.")
    ]
    dy = 370
    for t, d in tips:
        draw_bullet_dot(draw, 1415, dy + 10, radius=5, color=MP_BLUE)
        draw.text((1435, dy), t, font=FONT_BODY_BOLD, fill=MP_BLUE)
        dy += 28
        words = d.split()
        line = ""
        for w in words:
            if len(line + " " + w) > 30:
                draw.text((1435, dy), line, font=FONT_SMALL, fill=C_TEXT_MUTED)
                dy += 22
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1435, dy), line, font=FONT_SMALL, fill=C_TEXT_MUTED)
            dy += 35

    path = os.path.join(TEMP_DIR, "mod9_slide8.png")
    img.save(path)
    return path

def create_slide_9():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Checkout Pix", "Passo 6: Configurando o Checkout Transparente - Pix", "O meio de pagamento que mais cresce no Brasil com aprovacao em segundos", MP_BLUE)
    
    # Left container: Pix configuration fields
    draw.rounded_rectangle([80, 260, 920, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((115, 295), "Opcoes do Metodo Pix no WooCommerce", font=FONT_HEADING, fill=C_WHITE)
    
    pix_settings = [
        ("Ativar Checkout Pix:", "[X] Ativar Mercado Pago Pix como forma de pagamento", True),
        ("Titulo no Checkout:", "Pix - Aprovacao Instantanea", False),
        ("Descricao para o Cliente:", "Pague com Pix usando o QR Code ou Copia e Cola no app do seu banco.", False),
        ("Tempo de Expiracao do QR Code:", "30 minutos (Recomendado para evitar reservas eternas de estoque)", False),
        ("Desconto Promocional no Pix:", "5% de desconto automatico para quem pagar via Pix (Opcional)", False)
    ]
    
    for idx, (lbl, val, is_chk) in enumerate(pix_settings):
        py = 360 + idx * 115
        draw.text((115, py), lbl, font=FONT_SMALL_BOLD, fill=MELI_YELLOW)
        draw.rounded_rectangle([115, py + 26, 885, py + 75], radius=8, fill=(20, 28, 42), outline=(45, 55, 75), width=1)
        draw.text((135, py + 38), val, font=FONT_BODY, fill=C_TEXT_LIGHT)

    # Right container: Pix advantages & flow
    draw.rounded_rectangle([960, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=(16, 185, 129), width=2)
    draw.text((995, 295), "Fluxo do Pix Transparente", font=FONT_HEADING, fill=C_WHITE)
    
    pix_steps = [
        ("1. Cliente Clica em 'Finalizar Compra'", "Na pagina de checkout, escolhe Pix e conclui o pedido."),
        ("2. Geracao Instantanea do QR Code", "A tela exibe a imagem do QR Code dinamico e o botao 'Copiar codigo Pix'."),
        ("3. Pagamento no App Bancario", "O comprador abre o app do banco (Nubank, Inter, BB...) e efetua o pagamento."),
        ("4. Confirmacao por Webhook (IPN)", "O Banco Central confirma a transacao para o Mercado Pago, que avisa seu WooCommerce em menos de 3 segundos."),
        ("5. Status Atualizado Automaticamente", "O status do pedido muda de 'Pagamento Pendente' para 'Processando' sem nenhuma intervencao manual do lojista!")
    ]
    
    for idx, (step_t, step_d) in enumerate(pix_steps):
        sy = 360 + idx * 115
        draw.rounded_rectangle([995, sy, 1805, sy + 95], radius=10, fill=(24, 32, 47), outline=(45, 55, 75), width=1)
        draw.ellipse([1015, sy + 25, 1065, sy + 75], fill=(16, 185, 129))
        draw.text((1032, sy + 36), str(idx+1), font=FONT_HEADING, fill=C_WHITE)
        draw.text((1085, sy + 18), step_t, font=FONT_BODY_BOLD, fill=C_WHITE)
        draw.text((1085, sy + 52), step_d, font=FONT_SMALL, fill=C_TEXT_MUTED)

    path = os.path.join(TEMP_DIR, "mod9_slide9.png")
    img.save(path)
    return path

def create_slide_10():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Cartao de Credito", "Passo 7: Configurando Cartao de Credito e Parcelamento", "Taxas, maximo de parcelas e protecao antifraude para vendas no credito", MP_BLUE)
    
    # 2 Cards: Settings on left, Installment Strategy on right
    draw.rounded_rectangle([80, 260, 920, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((115, 295), "Parametros do Cartao de Credito", font=FONT_HEADING, fill=C_WHITE)
    
    card_settings = [
        ("Numero Maximo de Parcelas:", "Ate 12 parcelas"),
        ("Parcelas Sem Juros:", "Ex: 3x ou 6x sem juros (Configurado no Mercado Pago)"),
        ("Captura de Pagamento:", "Automatica (Imediata)"),
        ("Moeda de Cobranca:", "BRL (Real Brasileiro)"),
        ("Campos no Checkout:", "Numero do cartao, Nome impresso, Validade, CVV e CPF do titular")
    ]
    for idx, (lbl, val) in enumerate(card_settings):
        py = 360 + idx * 115
        draw.text((115, py), lbl, font=FONT_SMALL_BOLD, fill=MP_BLUE)
        draw.rounded_rectangle([115, py + 26, 885, py + 75], radius=8, fill=(20, 28, 42), outline=(45, 55, 75), width=1)
        draw.text((135, py + 38), val, font=FONT_BODY, fill=C_TEXT_LIGHT)

    # Right: Installment strategy
    draw.rounded_rectangle([960, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((995, 295), "Estrategia de Parcelamento e Juros", font=FONT_HEADING, fill=C_WHITE)
    
    strategies = [
        ("Quem paga os juros do parcelamento?", "No painel do Mercado Pago, voce escolhe se as tarifas de parcelamento serao absorvidas pela sua loja (oferecendo parcelas sem juros) ou repassadas ao cliente final.", (16, 185, 129)),
        ("Impacto na Conversao", "Oferecer 'Ate 3x ou 6x sem juros' aumenta a taxa de fechamento em produtos de ticket medio e alto (como roupas, calcados e eletronicos).", MP_BLUE),
        ("Protecao Antifraude Integrada", "O motor antifraude do Mercado Pago analisa mais de 5.000 variaveis por segundo para aprovar transacoes legitimas e barrar compras com cartoes clonados.", (239, 68, 68)),
        ("Recebimento dos Valores", "Mesmo que o cliente parcele em 12x, voce pode optar por receber o valor TOTAL da venda integralmente na sua conta conforme o plano de recebimento escolhido.", MELI_YELLOW)
    ]
    dy = 360
    for idx, (st_title, st_desc, st_col) in enumerate(strategies):
        draw.rounded_rectangle([995, dy, 1805, dy + 130], radius=12, fill=(24, 32, 47), outline=st_col, width=1)
        draw.rectangle([995, dy, 1007, dy + 130], fill=st_col)
        draw.text((1030, dy + 18), st_title, font=FONT_BODY_BOLD, fill=C_WHITE)
        
        words = st_desc.split()
        line = ""
        ty = dy + 52
        for w in words:
            if len(line + " " + w) > 55:
                draw.text((1030, ty), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                ty += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1030, ty), line, font=FONT_BODY, fill=C_TEXT_MUTED)
        dy += 150

    path = os.path.join(TEMP_DIR, "mod9_slide10.png")
    img.save(path)
    return path

def create_slide_11():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Marketplace", "Passo 8: Instalando o Plugin de Sincronizacao Mercado Livre", "Conectando o catalogo do WooCommerce com sua conta de vendedor Meli", MELI_YELLOW)
    
    # 3 Sequential setup steps
    steps = [
        ("1. Instalar o Plugin Oficial de Sincronizacao", "Painel WordPress > Plugins > Adicionar Novo > Buscar 'WooCommerce Mercado Libre Sync' ou plugin integrador oficial (como WooMeli).", MP_BLUE),
        ("2. Autenticacao OAuth 2.0 com o Mercado Livre", "Ao clicar em 'Conectar Conta', voce e redirecionado para a pagina segura do Mercado Livre para autorizar o acesso da sua loja.", MELI_YELLOW),
        ("3. Confirmacao de Vinculo de Loja", "Sua loja virtual recebe o token de integracao autorizado e pode comecar a exportar anuncios e sincronizar pedidos.", (16, 185, 129))
    ]
    
    for i, (title, desc, col) in enumerate(steps):
        y = 260 + i * 155
        draw.rounded_rectangle([80, y, 1840, y + 135], radius=16, fill=C_CARD_BG, outline=col, width=2)
        draw.ellipse([120, y + 25, 180, y + 85], fill=col)
        draw.text((142, y + 36), str(i+1), font=FONT_HEADING, fill=(20, 20, 20) if col == MELI_YELLOW else C_WHITE)
        
        draw.text((210, y + 25), title, font=FONT_HEADING, fill=C_WHITE)
        
        words = desc.split()
        line = ""
        dy = y + 70
        for w in words:
            if len(line + " " + w) > 90:
                draw.text((210, dy), line, font=FONT_BODY, fill=C_TEXT_MUTED)
                dy += 26
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((210, dy), line, font=FONT_BODY, fill=C_TEXT_MUTED)

    # Big info box below
    draw.rounded_rectangle([80, 750, 1840, 960], radius=16, fill=(24, 32, 47), outline=(45, 55, 75), width=1)
    draw.text((120, 775), "O que e transferido entre as plataformas?", font=FONT_HEADING, fill=C_WHITE)
    
    transfers = [
        ("Produtos e Fotos", "Titulo, descricao, precos, fotos e categorias sao publicados no Mercado Livre."),
        ("Estoque em Tempo Real", "Vendeu na loja? Diminui no Mercado Livre. Vendeu no Meli? Diminui na loja."),
        ("Pedidos e Envios", "Vendas no Mercado Livre entram como novos pedidos no painel do WooCommerce para despacho.")
    ]
    for idx, (tt, td) in enumerate(transfers):
        tx = 120 + idx * 560
        ty = 830
        draw.text((tx, ty), "● " + tt, font=FONT_BODY_BOLD, fill=MELI_YELLOW)
        draw.text((tx, ty + 30), td[:38], font=FONT_SMALL, fill=C_TEXT_LIGHT)
        draw.text((tx, ty + 55), td[38:], font=FONT_SMALL, fill=C_TEXT_LIGHT)

    path = os.path.join(TEMP_DIR, "mod9_slide11.png")
    img.save(path)
    return path

def create_slide_12():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Mapeamento", "Passo 9: Mapeamento de Categorias e Atributos", "Associando as categorias da sua loja a arvore oficial do Mercado Livre", MELI_YELLOW)
    
    # Comparison table / mapping mockup
    draw.rounded_rectangle([80, 260, 1840, 720], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((115, 290), "Tabela de Mapeamento de Categorias e Atributos", font=FONT_HEADING, fill=C_WHITE)
    
    headers = [("CATEGORIA NO WOOCOMMERCE", 120), ("ARVORE DE CATEGORIA NO MERCADO LIVRE", 650), ("STATUS DE VINCULO", 1450)]
    draw.rounded_rectangle([115, 340, 1805, 390], radius=6, fill=(20, 28, 42))
    for h_txt, hx in headers:
        draw.text((hx, 352), h_txt, font=FONT_BODY_BOLD, fill=MELI_YELLOW)
        
    mappings = [
        ("Moda > Camisetas Masculinas", "Calcados, Roupas e Bolsas > Camisas e Camisetas > Camisetas", "Mapeado com Sucesso", (16, 185, 129)),
        ("Calcados > Tenis Esportivos", "Calcados, Roupas e Bolsas > Tenis > Esportivos / Running", "Mapeado com Sucesso", (16, 185, 129)),
        ("Eletronicos > Fones Bluetooth", "Eletronicos, Audio e Video > Audio Portatil > Fones de Ouvido", "Mapeado com Sucesso", (16, 185, 129)),
        ("Acessorios > Relogios Inteligentes", "Smartwatches e Acessorios > Relogios Inteligentes", "Mapeado com Sucesso", (16, 185, 129))
    ]
    
    for idx, (woo_cat, meli_cat, st, col) in enumerate(mappings):
        my = 415 + idx * 75
        draw.rounded_rectangle([115, my, 1805, my + 60], radius=6, fill=(24, 32, 47), outline=(45, 55, 75), width=1)
        draw.text((120, my + 18), woo_cat, font=FONT_BODY_BOLD, fill=C_WHITE)
        draw.text((650, my + 18), meli_cat, font=FONT_BODY, fill=C_TEXT_LIGHT)
        draw_check_icon(draw, 1470, my + 30, radius=12, bg=col)
        draw.text((1495, my + 18), st, font=FONT_SMALL_BOLD, fill=col)

    # Bottom explanation
    draw.rounded_rectangle([80, 750, 1840, 960], radius=16, fill=(20, 28, 42), outline=MP_BLUE, width=1)
    draw.text((120, 775), "Por que o mapeamento correto e OBRIGATORIO?", font=FONT_HEADING, fill=MP_BLUE)
    draw.text((120, 815), "1. Cada categoria do Mercado Livre exige atributos tecnicos especificos (ex: Marca, Modelo, Genero, Cor, Tamanho).", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((120, 855), "2. Anuncios cadastrados na categoria errada perdem relevancia nas buscas e podem ser pausados pelos moderadores do marketplace.", font=FONT_BODY, fill=C_TEXT_LIGHT)
    draw.text((120, 895), "3. O mapeamento automatizado garante que novas variacoes adicionadas na loja subam com os campos certos.", font=FONT_BODY, fill=(16, 185, 129))

    path = os.path.join(TEMP_DIR, "mod9_slide12.png")
    img.save(path)
    return path

def create_slide_13():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Estoque e Precos", "Passo 10: Sincronizacao de Estoque e Precos", "Evitando o risco de furos de estoque e personalizando margens por canal", MP_BLUE)
    
    # 2 Big Cards: Estoque and Precos
    # Left: Estoque
    draw.rounded_rectangle([80, 260, 920, 980], radius=16, fill=C_CARD_BG, outline=(16, 185, 129), width=2)
    draw.text((115, 295), "Sincronizacao de Estoque em Tempo Real", font=FONT_HEADING, fill=C_WHITE)
    
    estoque_flow = [
        ("Cenario A: Venda na Loja Virtual", "Cliente compra a ultima unidade de um tenis na sua loja WooCommerce. O plugin envia uma requisicao imediata via API para o Mercado Livre pausando o anuncio automaticamente."),
        ("Cenario B: Venda no Mercado Livre", "Um cliente compra o mesmo produto no marketplace. O webhook avisa o WooCommerce e o estoque local diminui imediatamente, evitando que outro cliente compre o produto indisponivel."),
        ("Margem de Seguranca (Estoque Minimo):", "Voce pode configurar para que o anuncio no Mercado Livre seja pausado quando restarem 2 ou 3 unidades, reservando-as para a loja propria.")
    ]
    dy = 370
    for title, desc in estoque_flow:
        draw.text((115, dy), "● " + title, font=FONT_BODY_BOLD, fill=(16, 185, 129))
        dy += 28
        words = desc.split()
        line = ""
        for w in words:
            if len(line + " " + w) > 52:
                draw.text((135, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                dy += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((135, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
            dy += 45

    # Right: Precos
    draw.rounded_rectangle([960, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=MELI_YELLOW, width=2)
    draw.text((995, 295), "Regras de Preco Diferenciado", font=FONT_HEADING, fill=C_WHITE)
    
    preco_flow = [
        ("Por que cobrar precos diferentes?", "O Mercado Livre cobra uma comissao por venda (entre 11% e 16% conforme categoria e tipo de anuncio Classico vs Premium)."),
        ("Regra de Markup Automatico:", "O plugin permite definir um acrescimo percentual automatico no marketplace (ex: Preco no WooCommerce: R$ 100,00 -> Preco sincronizado no Meli: R$ 116,00)."),
        ("Preco Promocional na Loja Propria:", "Desta forma sua loja virtual sempre oferece o melhor preco ao consumidor final, incentivando as vendas diretas com maior margem de lucro!")
    ]
    dy = 370
    for title, desc in preco_flow:
        draw.text((995, dy), "● " + title, font=FONT_BODY_BOLD, fill=MELI_YELLOW)
        dy += 28
        words = desc.split()
        line = ""
        for w in words:
            if len(line + " " + w) > 52:
                draw.text((1015, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
                dy += 24
                line = w
            else:
                line = (line + " " + w).strip()
        if line:
            draw.text((1015, dy), line, font=FONT_BODY, fill=C_TEXT_LIGHT)
            dy += 45

    path = os.path.join(TEMP_DIR, "mod9_slide13.png")
    img.save(path)
    return path

def create_slide_14():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Operacao", "Passo 11: Gestao Unificada de Pedidos", "Como os pedidos de ambos os canais aparecem na mesma tela do WooCommerce", MP_BLUE)
    
    draw_browser_chrome(draw, 80, 260, 1840, 720, "WooCommerce - Gerenciar Pedidos", "https://meusite.com/wp-admin/edit.php?post_type=shop_order")
    
    # Table of unified orders
    draw.rectangle([80, 308, 1920, 980], fill=(245, 247, 250))
    draw.text((120, 340), "Pedidos de Venda", font=FONT_TITLE, fill=(30, 40, 55))
    
    # Table header
    draw.rounded_rectangle([120, 400, 1800, 450], radius=6, fill=(225, 235, 245))
    cols = [("PEDIDO", 140), ("DATA", 280), ("STATUS", 440), ("ORIGEM / CANAL", 680), ("COMPRADOR", 980), ("TOTAL", 1400), ("ACOES", 1620)]
    for c_name, cx in cols:
        draw.text((cx, 415), c_name, font=FONT_SMALL_BOLD, fill=(50, 70, 90))
        
    orders = [
        ("#1042", "Ha 5 minutos", "Processando", "Loja Virtual (Pix Mercado Pago)", "Carlos Silva - SP", "R$ 289,90", (16, 185, 129)),
        ("#1041", "Ha 32 minutos", "Processando", "Mercado Livre (Mercado Envios)", "Mariana Souza - RJ", "R$ 349,00", MELI_YELLOW),
        ("#1040", "Ha 1 hora", "Concluido", "Loja Virtual (Cartao Mercado Pago)", "Lucas Pereira - MG", "R$ 150,00", (16, 185, 129)),
        ("#1039", "Ha 3 horas", "Processando", "Mercado Livre (Mercado Envios)", "Roberto Lima - PR", "R$ 520,00", MELI_YELLOW),
        ("#1038", "Ontem", "Concluido", "Loja Virtual (Pix Mercado Pago)", "Fernanda Costa - RS", "R$ 89,90", (16, 185, 129))
    ]
    
    for idx, (p_num, dt, st, orig, comp, tot, col) in enumerate(orders):
        oy = 470 + idx * 75
        draw.rounded_rectangle([120, oy, 1800, oy + 65], radius=6, fill=C_WHITE, outline=(210, 220, 230), width=1)
        draw.text((140, oy + 20), p_num, font=FONT_BODY_BOLD, fill=(34, 113, 177))
        draw.text((280, oy + 22), dt, font=FONT_SMALL, fill=(100, 110, 125))
        
        # Status badge
        draw.rounded_rectangle([440, oy + 15, 620, oy + 50], radius=12, fill=(230, 248, 235))
        draw.text((465, oy + 20), st, font=FONT_SMALL_BOLD, fill=(16, 185, 129))
        
        # Channel origin badge
        chan_col = MELI_YELLOW if "Mercado Livre" in orig else MP_BLUE
        draw.text((680, oy + 20), orig, font=FONT_BODY_BOLD, fill=(180, 140, 0) if chan_col == MELI_YELLOW else MP_BLUE)
        
        draw.text((980, oy + 20), comp, font=FONT_BODY, fill=(40, 50, 65))
        draw.text((1400, oy + 20), tot, font=FONT_BODY_BOLD, fill=(30, 40, 55))
        
        draw.rounded_rectangle([1620, oy + 15, 1760, oy + 50], radius=6, fill=(34, 113, 177))
        draw.text((1645, oy + 22), "Ver Pedido", font=FONT_SMALL_BOLD, fill=C_WHITE)

    path = os.path.join(TEMP_DIR, "mod9_slide14.png")
    img.save(path)
    return path

def create_slide_15():
    img, draw = create_base_canvas()
    draw_header(draw, "Modulo 9 • Homologacao", "Checklist de Homologacao e Proximos Passos", "Validacao final da integracao de pagamentos e preparacao para o calculo de frete", MP_BLUE)
    
    draw.rounded_rectangle([80, 260, 1050, 980], radius=16, fill=C_CARD_BG, outline=C_CARD_BORDER, width=2)
    draw.text((115, 295), "Checklist de Pagamentos e Marketplace", font=FONT_HEADING, fill=C_WHITE)
    
    checklist = [
        "Plugin oficial do Mercado Pago instalado e ativo no WordPress",
        "Credenciais de API (Public Key e Access Token) inseridas",
        "Teste de compra com Pix simulado aprovado com sucesso",
        "Teste de compra com cartao em ambiente Sandbox aprovado",
        "Retorno automatico via Webhook funcionando (status muda para Processando)",
        "Plugin de integracao com Mercado Livre autenticado via OAuth",
        "Categorias mapeadas e sincronizacao de estoque testada",
        "Credenciais de Producao configuradas para receber vendas reais"
    ]
    for idx, item in enumerate(checklist):
        iy = 360 + idx * 72
        draw.rounded_rectangle([115, iy, 1015, iy + 60], radius=10, fill=(20, 30, 42), outline=(16, 185, 129), width=1)
        draw_check_icon(draw, 145, iy + 30, radius=12, bg=(16, 185, 129))
        draw.text((175, iy + 18), item, font=FONT_BODY, fill=C_TEXT_LIGHT)

    # Right: Next module preview
    draw.rounded_rectangle([1090, 260, 1840, 980], radius=16, fill=C_CARD_BG, outline=(14, 116, 144), width=2)
    draw.rounded_rectangle([1125, 290, 1805, 360], radius=10, fill=(14, 116, 144))
    draw.text((1150, 310), "PROXIMO PASSO: MODULO 10", font=FONT_HEADING, fill=C_WHITE)
    
    draw.text((1125, 390), "Integracao de Frete e Logistica", font=FONT_TITLE, fill=C_WHITE)
    draw.text((1125, 455), "Agora que os pagamentos estao prontos, vamos estruturar o calculo e envio das encomendas:", font=FONT_BODY, fill=C_TEXT_MUTED)
    
    next_topics = [
        ("Correios (PAC e SEDEX)", "Calculo em tempo real com base no peso e dimensoes dos produtos."),
        ("Melhor Envio & Transportadoras", "Cotacao simultanea com Jadlog, Loggi e descontos de ate 80% nas etiquetas."),
        ("Zonas de Entrega & Frete Gratis", "Criando faixas de CEP e regras de frete gratis atrativas para conversao.")
    ]
    ny = 525
    for nt, nd in next_topics:
        draw.rounded_rectangle([1125, ny, 1805, ny + 115], radius=12, fill=(24, 32, 47), outline=(45, 55, 75), width=1)
        draw.text((1150, ny + 18), "->  " + nt, font=FONT_HEADING, fill=(14, 116, 144))
        draw.text((1150, ny + 60), nd, font=FONT_BODY, fill=C_TEXT_LIGHT)
        ny += 135

    path = os.path.join(TEMP_DIR, "mod9_slide15.png")
    img.save(path)
    return path

def build_presentation_9():
    print("Generating slides for 9 plugin mercado livre.pptx...")
    slide_funcs = [
        create_slide_1, create_slide_2, create_slide_3, create_slide_4,
        create_slide_5, create_slide_6, create_slide_7, create_slide_8,
        create_slide_9, create_slide_10, create_slide_11, create_slide_12,
        create_slide_13, create_slide_14, create_slide_15
    ]
    
    prs = Presentation("6 wordpress.pptx")
    rIds = [prs.slides._sldIdLst[i].rId for i in range(len(prs.slides))]
    for rId in rIds:
        prs.part.drop_rel(rId)
    prs.slides._sldIdLst.clear()
    
    blank_layout = prs.slide_layouts[6]
    
    notes_scripts = [
        "Apresentacao do Modulo 9: Plugin Mercado Livre & Gateway Mercado Pago. Explicar a importancia estrategica de unificar vendas no marketplace com a loja propria.",
        "Apresentar os 3 pilares da estrategia omnichannel: Loja Virtual com controle total, Mercado Pago como processador confiavel e Mercado Livre como fonte massiva de novos clientes.",
        "Explicar a diferenca de funcoes entre Mercado Pago (transacoes financeiras e checkout) e Mercado Livre (catalogo e anuncios no marketplace).",
        "Guiar os alunos na instalacao do plugin oficial 'Mercado Pago payments for WooCommerce' pelo repositorio oficial do WordPress.",
        "Demonstrar o acesso ao Portal de Desenvolvedores do Mercado Pago (mercadopago.com.br/developers) e a criacao de uma nova aplicacao.",
        "Ensinar a diferenca crucial entre a Public Key (visivel publicamente no checkout) e o Access Token (chave secreta confidencial do servidor).",
        "Explicar a diferenca entre o Modo Sandbox (credenciais TEST- com cartoes virtuais) e o Modo Producao (credenciais APP_USR- para transacoes reais).",
        "Mostrar a tela de configuracao no WooCommerce onde o aluno deve colar as chaves obtidas e validar a conexao.",
        "Configuracao detalhada do Checkout Pix Transparente: tempo de expiracao do QR Code, chave Pix e funcionamento do webhook IPN de retorno instantaneo.",
        "Configuracao de Cartao de Credito: numero de parcelas, parcelas sem juros vs juros repassados e analise automatica do sistema antifraude.",
        "Instalacao do plugin integrador do Mercado Livre e autorizacao da conta de vendedor via OAuth 2.0 segura.",
        "Demonstrar a importancia do Mapeamento de Categorias: se uma camiseta for publicada na categoria errada do Mercado Livre, o anuncio perde visibilidade ou e rejeitado.",
        "Explicar a sincronizacao de estoque bidirecional e como configurar regras de precos diferentes para cobrir as comissoes do marketplace mantendo a loja propria mais competitiva.",
        "Mostrar a tela unificada de Pedidos no WooCommerce, onde as vendas da loja e do Mercado Livre sao visualizadas no mesmo lugar.",
        "Checklist final de homologacao de pagamentos e introducao do Modulo 10: Integracao de Frete e Logistica."
    ]
    
    for idx, func in enumerate(slide_funcs):
        img_path = func()
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(img_path, 0, 0, prs.slide_width, prs.slide_height)
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = notes_scripts[idx]
        print(f"  Slide {idx+1}/{len(slide_funcs)} generated.")

    out_file = "9 plugin mercado livre.pptx"
    prs.save(out_file)
    print(f"Presentation saved successfully as '{out_file}' ({os.path.getsize(out_file)} bytes).")

if __name__ == "__main__":
    build_presentation_9()
